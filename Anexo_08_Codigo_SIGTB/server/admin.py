"""Administración de cuentas: listado, alta, edición, baja y exportación."""

import csv
import io
from datetime import datetime
from functools import wraps

import bcrypt
from flask import Blueprint, current_app, jsonify, request, send_file, session
from openpyxl import Workbook

from server.authz import clear_sede_assignments_outside_empresa, user_has_permission
from server.db import get_empresas_connection, get_users_connection
from server.event_log import describe_usuario, log_evento
from server.onboarding import resolve_job_pending_notifications
from server.roles import (
    get_all_roles_catalog,
    is_operativo_director_job,
    is_pending_job,
    is_valid_managed_job,
    is_valid_managed_roll,
)
from server.validators import (
    is_letters_only,
    is_valid_email,
    normalize_email,
    sanitize_string,
    validate_password,
)
from server.limits import JOB, NAME

BCRYPT_ROUNDS = 12
EXPORT_COLUMNS = [
    "id_usuario",
    "usuario_login",
    "NAME_USER",
    "LAST_NAME_USER",
    "JOB",
    "ROLL",
    "empresa_id",
    "creation_date",
]

admin_bp = Blueprint("admin", __name__, url_prefix="/admin")


def _is_operativo_director_scope():
    return (
        session.get("ROLL") == "OPERATIVO"
        and is_operativo_director_job(session.get("JOB"))
        and user_has_permission("manage_users")
    )


def _actor_is_admin():
    return session.get("ROLL") == "ADMIN"


def _reject_non_admin_assigning_admin(payload_roll=None, target_roll=None):
    """El rango ADMIN solo lo asigna (o modifica) otro perfil Administrador."""
    if _actor_is_admin():
        return None
    if payload_roll == "ADMIN" or target_roll == "ADMIN":
        return jsonify(
            {
                "ok": False,
                "error": (
                    "El rango Administrador solo puede asignarse desde otro "
                    "perfil de Administrador. No requiere empresa, sede ni servicio."
                ),
            }
        ), 403
    return None


def _reject_admin_mutation(target_roll=None, payload_roll=None):
    """
    Solo Director Operativo: no puede modificar ni eliminar usuarios ADMIN
    bajo ningún concepto. El rol Administrador SÍ puede gestionar usuarios.
    """
    if not _is_operativo_director_scope():
        return None
    if target_roll == "ADMIN" or payload_roll == "ADMIN":
        return jsonify(
            {
                "ok": False,
                "error": (
                    "El Director Operativo no puede modificar ni eliminar "
                    "usuarios con rol Administrador."
                ),
            }
        ), 403
    return None


def admin_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if not session.get("id_usuario"):
            return jsonify({"ok": False, "error": "Debes iniciar sesión."}), 401
        if session.get("ROLL") != "ADMIN" and not (
            user_has_permission("manage_users")
            or user_has_permission("manage_asistencial_users")
        ):
            return jsonify(
                {
                    "ok": False,
                    "error": "Acceso denegado. Se requiere ADMIN, Director Operativo o Director asistencial.",
                }
            ), 403
        return view(*args, **kwargs)

    return wrapped


def _is_asistencial_director_scope():
    return (
        session.get("ROLL") == "ASISTENCIAL"
        and session.get("JOB") in ("Director", "Dirección")
        and user_has_permission("manage_asistencial_users")
        and not user_has_permission("manage_users")
    )


def _row_to_user(row):
    return {
        "id_usuario": row["id_usuario"],
        "usuario_login": row["usuario_login"],
        "NAME_USER": row["NAME_USER"],
        "LAST_NAME_USER": row["LAST_NAME_USER"],
        "JOB": row["JOB"],
        "ROLL": row["ROLL"],
        "creation_date": row["creation_date"],
        "empresa_id": row["empresa_id"] if "empresa_id" in row.keys() else None,
        "ID_Empresa": row["ID_Empresa"] if "ID_Empresa" in row.keys() else None,
        "ID_NIT": row["ID_NIT"] if "ID_NIT" in row.keys() else None,
    }


def _enrich_user_empresa(user_dict):
    empresa_id = user_dict.get("empresa_id")
    if not empresa_id:
        return user_dict
    with get_empresas_connection() as emp_conn:
        emp = emp_conn.execute(
            "SELECT ID_Empresa, ID_NIT FROM empresas WHERE id = ?",
            (empresa_id,),
        ).fetchone()
    if emp:
        user_dict["ID_Empresa"] = emp["ID_Empresa"]
        user_dict["ID_NIT"] = emp["ID_NIT"]
    return user_dict


def _empresa_exists(empresa_id):
    with get_empresas_connection() as emp_conn:
        return emp_conn.execute(
            "SELECT id FROM empresas WHERE id = ?",
            (empresa_id,),
        ).fetchone()


def _fetch_users(conn, *, empresa_id=None, only_asistencial=False, exclude_admin=False):
    if empresa_id is None:
        rows = conn.execute(
            """
            SELECT id_usuario, usuario_login, NAME_USER, LAST_NAME_USER,
                   JOB, ROLL, creation_date, empresa_id
            FROM usuarios
            ORDER BY id_usuario ASC
            """
        ).fetchall()
    elif exclude_admin:
        # Director Operativo: solo personal de su empresa (sin mezclar ADMIN).
        rows = conn.execute(
            """
            SELECT id_usuario, usuario_login, NAME_USER, LAST_NAME_USER,
                   JOB, ROLL, creation_date, empresa_id
            FROM usuarios
            WHERE empresa_id = ?
              AND ROLL != 'ADMIN'
            ORDER BY id_usuario ASC
            """,
            (empresa_id,),
        ).fetchall()
    else:
        rows = conn.execute(
            """
            SELECT id_usuario, usuario_login, NAME_USER, LAST_NAME_USER,
                   JOB, ROLL, creation_date, empresa_id
            FROM usuarios
            WHERE empresa_id = ? OR ROLL = 'ADMIN'
            ORDER BY id_usuario ASC
            """,
            (empresa_id,),
        ).fetchall()

    if only_asistencial:
        rows = [r for r in rows if r["ROLL"] == "ASISTENCIAL"]

    empresa_ids = {r["empresa_id"] for r in rows if r["empresa_id"]}
    empresas_map = {}
    if empresa_ids:
        with get_empresas_connection() as emp_conn:
            placeholders = ",".join("?" * len(empresa_ids))
            for e in emp_conn.execute(
                f"SELECT id, ID_Empresa, ID_NIT FROM empresas WHERE id IN ({placeholders})",
                tuple(empresa_ids),
            ).fetchall():
                empresas_map[e["id"]] = e

    result = []
    for row in rows:
        item = _row_to_user(row)
        emp = empresas_map.get(row["empresa_id"])
        if emp:
            item["ID_Empresa"] = emp["ID_Empresa"]
            item["ID_NIT"] = emp["ID_NIT"]
        result.append(item)
    return result


def _validate_user_payload(body, *, require_password=False, require_empresa=None):
    usuario_login = normalize_email(body.get("usuario_login"))
    password = body.get("password") if isinstance(body.get("password"), str) else ""
    name_user = sanitize_string(body.get("NAME_USER"), NAME)
    last_name_user = sanitize_string(body.get("LAST_NAME_USER"), NAME)
    job = sanitize_string(body.get("JOB"), JOB)
    roll = sanitize_string(body.get("ROLL"), 20).upper()
    creation_date = sanitize_string(body.get("creation_date"), 40)
    empresa_raw = body.get("empresa_id")
    if require_empresa is None:
        require_empresa = require_password

    field_errors = {}
    empresa_id = None
    if empresa_raw not in (None, ""):
        try:
            empresa_id = int(empresa_raw)
        except (TypeError, ValueError):
            field_errors["empresa_id"] = "empresa_id inválido."

    if not is_valid_email(usuario_login):
        field_errors["usuario_login"] = "Debe ser un correo válido."

    if require_password or password:
        password_check = validate_password(password)
        if not password_check["valid"]:
            field_errors["password"] = " ".join(password_check["errors"])

    if not is_letters_only(name_user):
        field_errors["NAME_USER"] = "Solo admite letras (sin números ni símbolos)."

    if not is_letters_only(last_name_user):
        field_errors["LAST_NAME_USER"] = "Solo admite letras (sin números ni símbolos)."

    if not is_valid_managed_roll(roll):
        field_errors["ROLL"] = "ROLL debe ser ASISTENCIAL, OPERATIVO o ADMIN."
    elif not is_valid_managed_job(roll, job):
        field_errors["JOB"] = f"JOB no válido para el rol {roll}."

    if roll != "ADMIN" and require_empresa and not empresa_id:
        field_errors["empresa_id"] = (
            "Debes asignar la empresa a la que pertenece el usuario."
        )

    if creation_date:
        try:
            datetime.strptime(creation_date, "%Y-%m-%d %H:%M:%S")
        except ValueError:
            try:
                datetime.strptime(creation_date, "%Y-%m-%d")
                creation_date = f"{creation_date} 00:00:00"
            except ValueError:
                field_errors["creation_date"] = (
                    "Usa formato YYYY-MM-DD o YYYY-MM-DD HH:MM:SS."
                )
    else:
        creation_date = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    payload = {
        "usuario_login": usuario_login,
        "password": password,
        "NAME_USER": name_user,
        "LAST_NAME_USER": last_name_user,
        "JOB": job,
        "ROLL": roll,
        "creation_date": creation_date,
        "empresa_id": empresa_id,
    }
    return payload, field_errors


@admin_bp.get("/roles")
@admin_required
def admin_roles():
    return jsonify({"ok": True, **get_all_roles_catalog()})


@admin_bp.get("/users")
@admin_required
def list_users():
    try:
        actor_empresa = session.get("empresa_id")
        only_asistencial = _is_asistencial_director_scope()
        operativo_director = _is_operativo_director_scope()
        scope_empresa = None
        if operativo_director or only_asistencial:
            scope_empresa = actor_empresa
        with get_users_connection() as conn:
            users = _fetch_users(
                conn,
                empresa_id=scope_empresa,
                only_asistencial=only_asistencial,
                exclude_admin=operativo_director,
            )
        return jsonify({"ok": True, "users": users, "total": len(users)})
    except Exception:
        current_app.logger.exception("Error al listar usuarios")
        return jsonify({"ok": False, "error": "Error interno del servidor."}), 500


@admin_bp.put("/users/<int:user_id>")
@admin_required
def update_user(user_id):
    try:
        body = request.get_json(silent=True) or {}
        payload, field_errors = _validate_user_payload(body)

        if field_errors:
            return jsonify(
                {
                    "ok": False,
                    "error": "Validación fallida. Corrige los campos indicados.",
                    "fieldErrors": field_errors,
                }
            ), 400

        with get_users_connection() as conn:
            # Director Operativo: solo usuarios de su empresa; no toca Administrador.
            if _is_operativo_director_scope():
                actor_empresa = session.get("empresa_id")
                payload["empresa_id"] = actor_empresa
                current_scope = conn.execute(
                    "SELECT empresa_id, ROLL FROM usuarios WHERE id_usuario = ?",
                    (user_id,),
                ).fetchone()
                blocked = _reject_admin_mutation(
                    target_roll=current_scope["ROLL"] if current_scope else None,
                    payload_roll=payload.get("ROLL"),
                )
                if blocked:
                    return blocked
                if (
                    current_scope
                    and current_scope["empresa_id"]
                    and current_scope["empresa_id"] != actor_empresa
                ):
                    return jsonify(
                        {
                            "ok": False,
                            "error": "No puedes editar usuarios de otra empresa.",
                        }
                    ), 403

            # Director asistencial: solo personal ASISTENCIAL de su empresa.
            if _is_asistencial_director_scope():
                actor_empresa = session.get("empresa_id")
                payload["empresa_id"] = actor_empresa
                payload["ROLL"] = "ASISTENCIAL"
                target = conn.execute(
                    "SELECT empresa_id, ROLL FROM usuarios WHERE id_usuario = ?",
                    (user_id,),
                ).fetchone()
                if not target or target["ROLL"] != "ASISTENCIAL":
                    return jsonify(
                        {
                            "ok": False,
                            "error": "Solo puedes gestionar personal asistencial.",
                        }
                    ), 403
                if target["empresa_id"] != actor_empresa:
                    return jsonify(
                        {
                            "ok": False,
                            "error": "No puedes editar usuarios de otra empresa.",
                        }
                    ), 403

            if payload.get("empresa_id"):
                if not _empresa_exists(payload["empresa_id"]):
                    return jsonify(
                        {
                            "ok": False,
                            "error": "La empresa asignada no existe.",
                            "fieldErrors": {"empresa_id": "Empresa inválida."},
                        }
                    ), 400

            current = conn.execute(
                """
                SELECT id_usuario, ROLL, empresa_id
                FROM usuarios WHERE id_usuario = ?
                """,
                (user_id,),
            ).fetchone()
            if current is None:
                return jsonify({"ok": False, "error": "Usuario no encontrado."}), 404

            blocked_admin = _reject_non_admin_assigning_admin(
                payload_roll=payload.get("ROLL"),
                target_roll=current["ROLL"],
            )
            if blocked_admin:
                return blocked_admin

            duplicate = conn.execute(
                """
                SELECT id_usuario FROM usuarios
                WHERE usuario_login = ? COLLATE NOCASE AND id_usuario != ?
                """,
                (payload["usuario_login"], user_id),
            ).fetchone()
            if duplicate:
                return jsonify(
                    {
                        "ok": False,
                        "error": "Ya existe otro usuario con ese correo.",
                        "fieldErrors": {
                            "usuario_login": "Este correo ya está registrado."
                        },
                    }
                ), 409

            # No dejar el sistema sin ningún ADMIN.
            if current["ROLL"] == "ADMIN" and payload["ROLL"] != "ADMIN":
                admin_count = conn.execute(
                    "SELECT COUNT(*) AS total FROM usuarios WHERE ROLL = 'ADMIN'"
                ).fetchone()["total"]
                if admin_count <= 1:
                    return jsonify(
                        {
                            "ok": False,
                            "error": "Debe existir al menos un usuario ADMIN.",
                            "fieldErrors": {
                                "ROLL": "No puedes quitar el último administrador."
                            },
                        }
                    ), 400

            if payload["password"]:
                password_hash = bcrypt.hashpw(
                    payload["password"].encode("utf-8"),
                    bcrypt.gensalt(rounds=BCRYPT_ROUNDS),
                ).decode("utf-8")
                conn.execute(
                    """
                    UPDATE usuarios
                    SET usuario_login = ?, password_hash = ?, NAME_USER = ?,
                        LAST_NAME_USER = ?, JOB = ?, ROLL = ?, creation_date = ?,
                        empresa_id = ?
                    WHERE id_usuario = ?
                    """,
                    (
                        payload["usuario_login"],
                        password_hash,
                        payload["NAME_USER"],
                        payload["LAST_NAME_USER"],
                        payload["JOB"],
                        payload["ROLL"],
                        payload["creation_date"],
                        payload["empresa_id"],
                        user_id,
                    ),
                )
            else:
                conn.execute(
                    """
                    UPDATE usuarios
                    SET usuario_login = ?, NAME_USER = ?, LAST_NAME_USER = ?,
                        JOB = ?, ROLL = ?, creation_date = ?, empresa_id = ?
                    WHERE id_usuario = ?
                    """,
                    (
                        payload["usuario_login"],
                        payload["NAME_USER"],
                        payload["LAST_NAME_USER"],
                        payload["JOB"],
                        payload["ROLL"],
                        payload["creation_date"],
                        payload["empresa_id"],
                        user_id,
                    ),
                )

            # Si cambió de empresa (o quedó sin ella), quitar sedes de otras empresas.
            # ADMIN no requiere sede ni servicio: se limpian asignaciones.
            prev_empresa = current["empresa_id"]
            new_empresa = payload.get("empresa_id")
            if payload["ROLL"] == "ADMIN":
                clear_sede_assignments_outside_empresa(conn, user_id, None)
            elif prev_empresa != new_empresa:
                clear_sede_assignments_outside_empresa(conn, user_id, new_empresa)

            if not is_pending_job(payload["JOB"]):
                resolve_job_pending_notifications(conn, user_id)

            conn.commit()

            updated = conn.execute(
                """
                SELECT id_usuario, usuario_login, NAME_USER, LAST_NAME_USER,
                       JOB, ROLL, creation_date, empresa_id
                FROM usuarios
                WHERE id_usuario = ?
                """,
                (user_id,),
            ).fetchone()

            # Si el usuario se edita a sí mismo, refrescar sesión.
            if session.get("id_usuario") == user_id:
                session["usuario_login"] = updated["usuario_login"]
                session["ROLL"] = updated["ROLL"]
                session["JOB"] = updated["JOB"]
                session["NAME_USER"] = updated["NAME_USER"]
                session["empresa_id"] = updated["empresa_id"]

        label = describe_usuario(updated)
        if prev_empresa and prev_empresa != new_empresa:
            log_evento(
                empresa_id=prev_empresa,
                accion=f"Baja de usuario {label}.",
            )
        if new_empresa and prev_empresa != new_empresa:
            log_evento(
                empresa_id=new_empresa,
                accion=f"Alta de usuario {label}.",
            )

        return jsonify(
            {
                "ok": True,
                "message": "Usuario actualizado correctamente.",
                "user": _enrich_user_empresa(_row_to_user(updated)),
            }
        )
    except Exception:
        current_app.logger.exception("Error interno al actualizar el usuario")
        return jsonify(
            {"ok": False, "error": "Error interno del servidor."}
        ), 500


@admin_bp.post("/users")
@admin_required
def create_user():
    try:
        body = request.get_json(silent=True) or {}
        payload, field_errors = _validate_user_payload(body, require_password=True)
        if field_errors:
            return jsonify(
                {
                    "ok": False,
                    "error": "Validación fallida. Corrige los campos indicados.",
                    "fieldErrors": field_errors,
                }
            ), 400

        blocked_admin = _reject_non_admin_assigning_admin(payload_roll=payload.get("ROLL"))
        if blocked_admin:
            return blocked_admin

        if (
            session.get("ROLL") == "OPERATIVO"
            and is_operativo_director_job(session.get("JOB"))
        ):
            payload["empresa_id"] = session.get("empresa_id")
            blocked = _reject_admin_mutation(payload_roll=payload.get("ROLL"))
            if blocked:
                return blocked

        if _is_asistencial_director_scope():
            payload["empresa_id"] = session.get("empresa_id")
            payload["ROLL"] = "ASISTENCIAL"
            if payload.get("JOB") not in ("Director", "Coordinador", "Líder", "Dirección"):
                return jsonify(
                    {
                        "ok": False,
                        "error": "Solo puedes asignar cargos asistenciales (Director, Coordinador, Líder).",
                        "fieldErrors": {"JOB": "Cargo no permitido."},
                    }
                ), 400

        with get_users_connection() as conn:
            if payload.get("empresa_id"):
                if not _empresa_exists(payload["empresa_id"]):
                    return jsonify(
                        {
                            "ok": False,
                            "error": "La empresa asignada no existe.",
                            "fieldErrors": {"empresa_id": "Empresa inválida."},
                        }
                    ), 400

            existing = conn.execute(
                "SELECT id_usuario FROM usuarios WHERE usuario_login = ? COLLATE NOCASE",
                (payload["usuario_login"],),
            ).fetchone()
            if existing:
                return jsonify(
                    {
                        "ok": False,
                        "error": "Ya existe un usuario con ese correo.",
                        "fieldErrors": {
                            "usuario_login": "Este correo ya está registrado."
                        },
                    }
                ), 409

            password_hash = bcrypt.hashpw(
                payload["password"].encode("utf-8"),
                bcrypt.gensalt(rounds=BCRYPT_ROUNDS),
            ).decode("utf-8")

            cur = conn.execute(
                """
                INSERT INTO usuarios (
                    usuario_login, password_hash, creation_date,
                    NAME_USER, LAST_NAME_USER, JOB, ROLL, empresa_id
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    payload["usuario_login"],
                    password_hash,
                    payload["creation_date"],
                    payload["NAME_USER"],
                    payload["LAST_NAME_USER"],
                    payload["JOB"],
                    payload["ROLL"],
                    payload["empresa_id"],
                ),
            )
            conn.commit()
            created = conn.execute(
                """
                SELECT id_usuario, usuario_login, NAME_USER, LAST_NAME_USER,
                       JOB, ROLL, creation_date, empresa_id
                FROM usuarios
                WHERE id_usuario = ?
                """,
                (cur.lastrowid,),
            ).fetchone()

        if created and created["empresa_id"]:
            log_evento(
                empresa_id=created["empresa_id"],
                accion=f"Alta de usuario {describe_usuario(created)}.",
            )

        return jsonify(
            {
                "ok": True,
                "message": "Usuario creado correctamente.",
                "user": _enrich_user_empresa(_row_to_user(created)),
            }
        ), 201
    except Exception:
        current_app.logger.exception("Error interno al crear el usuario")
        return jsonify(
            {"ok": False, "error": "Error interno del servidor."}
        ), 500


@admin_bp.delete("/users/<int:user_id>")
@admin_required
def delete_user(user_id):
    try:
        if session.get("id_usuario") == user_id:
            return jsonify(
                {"ok": False, "error": "No puedes eliminar tu propia cuenta."}
            ), 400

        with get_users_connection() as conn:
            current = conn.execute(
                """
                SELECT id_usuario, usuario_login, NAME_USER, LAST_NAME_USER,
                       JOB, ROLL, empresa_id
                FROM usuarios
                WHERE id_usuario = ?
                """,
                (user_id,),
            ).fetchone()
            if current is None:
                return jsonify({"ok": False, "error": "Usuario no encontrado."}), 404

            if (
                session.get("ROLL") == "OPERATIVO"
                and is_operativo_director_job(session.get("JOB"))
            ):
                blocked = _reject_admin_mutation(target_roll=current["ROLL"])
                if blocked:
                    return blocked
                if current["empresa_id"] != session.get("empresa_id"):
                    return jsonify(
                        {"ok": False, "error": "No puedes eliminar usuarios de otra empresa."}
                    ), 403

            if _is_asistencial_director_scope():
                if current["ROLL"] != "ASISTENCIAL":
                    return jsonify(
                        {"ok": False, "error": "Solo puedes eliminar personal asistencial."}
                    ), 403
                if current["empresa_id"] != session.get("empresa_id"):
                    return jsonify(
                        {"ok": False, "error": "No puedes eliminar usuarios de otra empresa."}
                    ), 403

            if current["ROLL"] == "ADMIN":
                admin_count = conn.execute(
                    "SELECT COUNT(*) AS total FROM usuarios WHERE ROLL = 'ADMIN'"
                ).fetchone()["total"]
                if admin_count <= 1:
                    return jsonify(
                        {"ok": False, "error": "No puedes eliminar el último ADMIN."}
                    ), 400

            conn.execute("DELETE FROM usuarios WHERE id_usuario = ?", (user_id,))
            conn.commit()

        if current and current["empresa_id"]:
            log_evento(
                empresa_id=current["empresa_id"],
                accion=f"Baja de usuario {describe_usuario(current)}.",
            )

        return jsonify({"ok": True, "message": "Usuario eliminado correctamente."})
    except Exception:
        current_app.logger.exception("Error interno al eliminar el usuario")
        return jsonify(
            {"ok": False, "error": "Error interno del servidor."}
        ), 500


@admin_bp.get("/users/export.csv")
@admin_required
def export_users_csv():
    try:
        actor_empresa = session.get("empresa_id")
        only_asistencial = _is_asistencial_director_scope()
        operativo_director = _is_operativo_director_scope()
        scope_empresa = None
        if operativo_director or only_asistencial:
            scope_empresa = actor_empresa
        with get_users_connection() as conn:
            users = _fetch_users(
                conn,
                empresa_id=scope_empresa,
                only_asistencial=only_asistencial,
                exclude_admin=operativo_director,
            )

        buffer = io.StringIO()
        writer = csv.DictWriter(buffer, fieldnames=EXPORT_COLUMNS, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(users)

        data = io.BytesIO(buffer.getvalue().encode("utf-8-sig"))
        data.seek(0)
        filename = f"usuarios_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
        return send_file(
            data,
            mimetype="text/csv",
            as_attachment=True,
            download_name=filename,
        )
    except Exception:
        current_app.logger.exception("Error al exportar CSV de usuarios")
        return jsonify({"ok": False, "error": "Error interno del servidor."}), 500


@admin_bp.get("/users/export.xlsx")
@admin_required
def export_users_xlsx():
    try:
        actor_empresa = session.get("empresa_id")
        only_asistencial = _is_asistencial_director_scope()
        operativo_director = _is_operativo_director_scope()
        scope_empresa = None
        if operativo_director or only_asistencial:
            scope_empresa = actor_empresa
        with get_users_connection() as conn:
            users = _fetch_users(
                conn,
                empresa_id=scope_empresa,
                only_asistencial=only_asistencial,
                exclude_admin=operativo_director,
            )

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "usuarios"
        sheet.append(EXPORT_COLUMNS)
        for user in users:
            sheet.append([user.get(col, "") for col in EXPORT_COLUMNS])

        data = io.BytesIO()
        workbook.save(data)
        data.seek(0)
        filename = f"usuarios_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        return send_file(
            data,
            mimetype=(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            ),
            as_attachment=True,
            download_name=filename,
        )
    except Exception:
        current_app.logger.exception("Error al exportar XLSX de usuarios")
        return jsonify({"ok": False, "error": "Error interno del servidor."}), 500


def _backup_admin_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if not session.get("id_usuario"):
            return jsonify({"ok": False, "error": "Debes iniciar sesión."}), 401
        if session.get("ROLL") != "ADMIN":
            return jsonify(
                {
                    "ok": False,
                    "error": "Solo el Administrador puede gestionar respaldos.",
                }
            ), 403
        return view(*args, **kwargs)

    return wrapped


@admin_bp.get("/backups")
@_backup_admin_required
def list_backups():
    from server.backup import BACKUP_DIR, KEEP_SNAPSHOTS, list_snapshots, offsite_dir

    return jsonify(
        {
            "ok": True,
            "backups": list_snapshots(),
            "keep": KEEP_SNAPSHOTS,
            "local_dir": str(BACKUP_DIR),
            "offsite_dir": str(offsite_dir()),
        }
    )


@admin_bp.post("/backups")
@_backup_admin_required
def create_backup_now():
    from server.backup import create_snapshot

    result = create_snapshot("manual")
    status = 200 if result.get("ok") else 500
    return jsonify(result), status


@admin_bp.post("/backups/<snapshot_id>/restore")
@_backup_admin_required
def restore_backup(snapshot_id):
    from server.backup import restore_snapshot

    safe_id = sanitize_string(snapshot_id or "")
    if not safe_id or "/" in safe_id or "\\" in safe_id or ".." in safe_id:
        return jsonify({"ok": False, "error": "Identificador de respaldo no válido."}), 400
    result = restore_snapshot(safe_id)
    status = 200 if result.get("ok") else 400
    return jsonify(result), status
