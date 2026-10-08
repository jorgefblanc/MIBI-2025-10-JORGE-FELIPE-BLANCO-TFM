"""CRUD y persistencia de ejecución de mantenimientos (preventivo/correctivo/visita)."""

from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path

from server.frecuencia_pm import (
    FUENTES_EJECUCION,
    TIPOS_EJECUCION,
    attach_resumen_ejecucion,
    estructura_frecuencia_equipo,
    llave_compuesta,
    lookup_resumen_llave,
    normalize_codigo,
    parse_date,
    resumen_ejecucion_por_llaves,
)
from server.limits import CODE, COMMENT, IMPORT_ROWS, LABEL, over_limit
from server.validators import sanitize_string

SCHEMA_PATH = Path(__file__).resolve().parent.parent / "sql" / "schema_ejecucion_mantenimientos.sql"
SCHEMA_SQL = SCHEMA_PATH.read_text(encoding="utf-8")

PAGE_SIZES = (10, 25, 50)
DEFAULT_PAGE_SIZE = 25
MAX_DURACION_HORAS = 720.0
FOSCAGIB_EJECUCION_CANDIDATES = (
    Path(__file__).resolve().parent.parent / "data" / "mis_datos_foscagib.xlsx",
)


def ensure_ejecucion_schema(emp_conn) -> None:
    emp_conn.executescript(SCHEMA_SQL)
    cols = {
        (row["name"] if hasattr(row, "keys") else row[1])
        for row in emp_conn.execute("PRAGMA table_info(ejecucion_mantenimientos)").fetchall()
    }
    if "fecha_atencion" not in cols:
        emp_conn.execute("ALTER TABLE ejecucion_mantenimientos ADD COLUMN fecha_atencion TEXT")
    if "fecha_cierre" not in cols:
        emp_conn.execute("ALTER TABLE ejecucion_mantenimientos ADD COLUMN fecha_cierre TEXT")
    tables = {
        row["name"] if hasattr(row, "keys") else row[0]
        for row in emp_conn.execute("SELECT name FROM sqlite_master WHERE type='table'").fetchall()
    }
    if "inventario_equipos" not in tables:
        return
    emp_conn.executescript(
        """
        CREATE INDEX IF NOT EXISTS idx_inv_num_biomedica_norm
            ON inventario_equipos (
                REPLACE(REPLACE(UPPER(TRIM(COALESCE(num_biomedica, ''))), ' ', ''), '-', '')
            );
        CREATE INDEX IF NOT EXISTS idx_inv_codigo_activo_norm
            ON inventario_equipos (
                REPLACE(REPLACE(UPPER(TRIM(COALESCE(codigo_activo, ''))), ' ', ''), '-', '')
            );
        """
    )


def clamp_page(page, size) -> tuple[int, int, int]:
    try:
        size_n = int(size)
    except (TypeError, ValueError):
        size_n = DEFAULT_PAGE_SIZE
    if size_n not in PAGE_SIZES:
        size_n = DEFAULT_PAGE_SIZE
    try:
        page_n = int(page)
    except (TypeError, ValueError):
        page_n = 1
    page_n = max(1, page_n)
    return page_n, size_n, (page_n - 1) * size_n


def _tipo(value) -> str | None:
    text = sanitize_string(str(value or ""), 40).casefold()
    if text in TIPOS_EJECUCION:
        return text
    aliases = {
        "pm": "preventivo",
        "mp": "preventivo",
        "mantenimiento preventivo": "preventivo",
        "mc": "correctivo",
        "mantenimiento correctivo": "correctivo",
        "visita tecnica": "visita",
        "visita técnica": "visita",
        "servicio": "visita",
        "otro": "visita",
        "calibracion": "visita",
        "calibración": "visita",
    }
    return aliases.get(text)


def _fuente(value, default="manual") -> str:
    text = sanitize_string(str(value or default), 40).casefold()
    return text if text in FUENTES_EJECUCION else default


def parse_datetime(value) -> datetime | None:
    """Acepta date, datetime o ISO (YYYY-MM-DD[THH:MM[:SS]])."""
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    text = str(value).strip().replace("Z", "")
    if "T" not in text[:19]:
        text = text.replace(" ", "T", 1)
    for size, fmt in (
        (19, "%Y-%m-%dT%H:%M:%S"),
        (16, "%Y-%m-%dT%H:%M"),
        (10, "%Y-%m-%d"),
    ):
        chunk = text[:size]
        try:
            return datetime.strptime(chunk, fmt)
        except ValueError:
            continue
    parsed = parse_date(value)
    if parsed:
        return datetime(parsed.year, parsed.month, parsed.day)
    return None


def duracion_desde_intervalo(atencion, cierre):
    """Horas = (fecha de cierre − fecha de atención). None si no es calculable."""
    start = parse_datetime(atencion)
    end = parse_datetime(cierre)
    if not start or not end:
        return None
    delta = (end - start).total_seconds() / 3600.0
    if delta <= 0 or delta > MAX_DURACION_HORAS:
        return None
    return round(delta, 2)


def _duracion(value):
    if value in (None, ""):
        return None
    try:
        hours = float(value)
    except (TypeError, ValueError):
        return None
    if hours < 0 or hours > MAX_DURACION_HORAS:
        return None
    return round(hours, 2)


def _iso_datetime(value: datetime | None) -> str | None:
    if value is None:
        return None
    if value.hour or value.minute or value.second:
        return value.strftime("%Y-%m-%dT%H:%M:%S")
    return value.strftime("%Y-%m-%d")


def row_payload(row) -> dict:
    data = dict(row)
    for key in ("fecha_ejecucion", "fecha_atencion", "fecha_cierre", "created_at", "updated_at"):
        if data.get(key) is not None:
            text = str(data[key])
            if key in {"fecha_ejecucion"}:
                data[key] = text[:10]
            elif key in {"fecha_atencion", "fecha_cierre"}:
                data[key] = text[:19]
            else:
                data[key] = text[:19]
    return data


def validate_manual_payload(body: dict) -> tuple[dict | None, str | None]:
    body = body or {}
    bio_raw = sanitize_string(body.get("codigo_biomedica"), CODE)
    act_raw = sanitize_string(body.get("codigo_activo"), CODE)
    if over_limit(str(body.get("codigo_biomedica") or ""), CODE) or over_limit(
        str(body.get("codigo_activo") or ""), CODE
    ):
        return None, "Código Biomédica o Código de Activo supera el límite."
    bio_norm, act_norm = llave_compuesta(bio_raw, act_raw)
    if not bio_norm and not act_norm:
        return None, "Indique Código Biomédica o Código de Activo."
    tipo = _tipo(body.get("tipo_mantenimiento"))
    if not tipo:
        return None, "Tipo debe ser preventivo, correctivo o visita."
    fecha = parse_date(body.get("fecha_ejecucion"))
    atencion = parse_datetime(body.get("fecha_atencion"))
    cierre = parse_datetime(body.get("fecha_cierre"))
    if body.get("fecha_atencion") not in (None, "") and not atencion:
        return None, "Fecha de atención inválida."
    if body.get("fecha_cierre") not in (None, "") and not cierre:
        return None, "Fecha de cierre inválida."
    if bool(atencion) != bool(cierre):
        return None, "Indique fecha de atención y fecha de cierre para calcular la duración."
    computed = duracion_desde_intervalo(atencion, cierre)
    if atencion and cierre and computed is None:
        return None, "La fecha de cierre debe ser posterior a la de atención (máximo 720 h)."
    if not fecha and atencion:
        fecha = atencion.date()
    if not fecha:
        return None, "Fecha de ejecución inválida."
    if fecha.year < 1990 or fecha.year > 2100:
        return None, "Fecha de ejecución fuera de rango."
    obs_raw = body.get("observaciones") or ""
    if over_limit(str(obs_raw), COMMENT):
        return None, "Observaciones supera el límite de 500 caracteres."
    proveedor_raw = body.get("proveedor") or ""
    if over_limit(str(proveedor_raw), LABEL):
        return None, "Proveedor supera el límite."
    if computed is not None:
        duracion = computed
    else:
        duracion = _duracion(body.get("duracion_horas"))
        if body.get("duracion_horas") not in (None, "") and duracion is None:
            return None, "Duración en horas inválida."
    try:
        empresa_id = int(body.get("empresa_id"))
    except (TypeError, ValueError):
        return None, "empresa_id inválido."
    sede_id = body.get("sede_id")
    servicio_id = body.get("servicio_id")
    try:
        sede_id = int(sede_id) if sede_id not in (None, "") else None
    except (TypeError, ValueError):
        return None, "sede_id inválido."
    try:
        servicio_id = int(servicio_id) if servicio_id not in (None, "") else None
    except (TypeError, ValueError):
        return None, "servicio_id inválido."
    return {
        "codigo_biomedica": bio_raw or bio_norm,
        "codigo_activo": act_raw if act_raw else (body.get("codigo_activo") or ""),
        "codigo_biomedica_norm": bio_norm,
        "codigo_activo_norm": act_norm,
        "empresa_id": empresa_id,
        "sede_id": sede_id,
        "servicio_id": servicio_id,
        "tipo_mantenimiento": tipo,
        "fecha_ejecucion": fecha.isoformat(),
        "fecha_atencion": _iso_datetime(atencion),
        "fecha_cierre": _iso_datetime(cierre),
        "proveedor": sanitize_string(proveedor_raw, LABEL) or None,
        "duracion_horas": duracion,
        "observaciones": sanitize_string(str(obs_raw), COMMENT) or None,
        "fuente": _fuente(body.get("fuente"), "manual"),
    }, None


def insert_ejecucion(conn, payload: dict) -> dict:
    cur = conn.execute(
        """
        INSERT INTO ejecucion_mantenimientos (
            codigo_biomedica, codigo_activo, codigo_biomedica_norm, codigo_activo_norm,
            empresa_id, sede_id, servicio_id, tipo_mantenimiento, fecha_ejecucion,
            fecha_atencion, fecha_cierre,
            proveedor, duracion_horas, observaciones, fuente, created_at, updated_at
        ) VALUES (
            ?, ?, ?, ?,
            ?, ?, ?, ?, ?,
            ?, ?,
            ?, ?, ?, ?, datetime('now'), datetime('now')
        )
        """,
        (
            payload["codigo_biomedica"],
            payload["codigo_activo"] or "",
            payload["codigo_biomedica_norm"],
            payload["codigo_activo_norm"] or "",
            payload["empresa_id"],
            payload.get("sede_id"),
            payload.get("servicio_id"),
            payload["tipo_mantenimiento"],
            payload.get("fecha_ejecucion"),
            payload.get("fecha_atencion"),
            payload.get("fecha_cierre"),
            payload.get("proveedor"),
            payload.get("duracion_horas"),
            payload.get("observaciones"),
            payload.get("fuente") or "manual",
        ),
    )
    row = conn.execute(
        "SELECT * FROM ejecucion_mantenimientos WHERE id = ?",
        (cur.lastrowid,),
    ).fetchone()
    return row_payload(row)


def list_ejecucion(
    conn,
    empresa_id: int,
    *,
    sede_id=None,
    servicio_id=None,
    tipo=None,
    q=None,
    sede_ids: set[int] | None = None,
    page=1,
    page_size=DEFAULT_PAGE_SIZE,
) -> dict:
    page_n, size_n, offset = clamp_page(page, page_size)
    clauses = ["em.empresa_id = ?"]
    params: list = [int(empresa_id)]
    if sede_id not in (None, ""):
        clauses.append("em.sede_id = ?")
        params.append(int(sede_id))
    if servicio_id not in (None, ""):
        clauses.append("em.servicio_id = ?")
        params.append(int(servicio_id))
    tipo_n = _tipo(tipo) if tipo not in (None, "") else None
    if tipo_n:
        clauses.append("em.tipo_mantenimiento = ?")
        params.append(tipo_n)
    if sede_ids:
        ids = sorted({int(x) for x in sede_ids})
        placeholders = ",".join("?" * len(ids))
        clauses.append(f"(em.sede_id IS NULL OR em.sede_id IN ({placeholders}))")
        params.extend(ids)
    q_norm = normalize_codigo(q)
    if q_norm:
        clauses.append(
            "(em.codigo_biomedica_norm LIKE ? OR em.codigo_activo_norm LIKE ? OR IFNULL(em.proveedor, '') LIKE ?)"
        )
        like = f"%{q_norm}%"
        params.extend((like, like, f"%{sanitize_string(str(q), 80)}%"))
    where = " AND ".join(clauses)
    total = conn.execute(
        f"SELECT COUNT(*) AS n FROM ejecucion_mantenimientos em WHERE {where}",
        params,
    ).fetchone()["n"]
    rows = conn.execute(
        f"""
        SELECT em.*, s.name_sede, s.ID_sede, srv.name_servicio, srv.ID_servicio
        FROM ejecucion_mantenimientos em
        LEFT JOIN sedes s ON s.id = em.sede_id
        LEFT JOIN servicios srv ON srv.id = em.servicio_id
        WHERE {where}
        ORDER BY em.fecha_ejecucion DESC, em.id DESC
        LIMIT ? OFFSET ?
        """,
        (*params, size_n, offset),
    ).fetchall()
    return {
        "ok": True,
        "total": int(total or 0),
        "page": page_n,
        "page_size": size_n,
        "items": [row_payload(r) for r in rows],
    }


def list_ejecucion_por_llave(
    conn,
    empresa_id: int,
    codigo_biomedica,
    codigo_activo,
    *,
    page=1,
    page_size=DEFAULT_PAGE_SIZE,
) -> dict:
    page_n, size_n, offset = clamp_page(page, page_size)
    bio, act = llave_compuesta(codigo_biomedica, codigo_activo)
    if not bio and not act:
        return {"ok": True, "total": 0, "page": page_n, "page_size": size_n, "items": []}
    clauses = ["em.empresa_id = ?"]
    params: list = [int(empresa_id)]
    if bio and act:
        clauses.append(
            """(
                (em.codigo_biomedica_norm = ? AND (em.codigo_activo_norm = ? OR em.codigo_activo_norm = ''))
                OR (em.codigo_activo_norm = ? AND (em.codigo_biomedica_norm = ? OR em.codigo_biomedica_norm = ''))
            )"""
        )
        params.extend([bio, act, act, bio])
    elif bio:
        clauses.append("em.codigo_biomedica_norm = ?")
        params.append(bio)
    else:
        clauses.append("em.codigo_activo_norm = ?")
        params.append(act)
    where = " AND ".join(clauses)
    total = conn.execute(
        f"SELECT COUNT(*) AS n FROM ejecucion_mantenimientos em WHERE {where}",
        params,
    ).fetchone()["n"]
    rows = conn.execute(
        f"""
        SELECT em.*, s.name_sede, s.ID_sede, srv.name_servicio, srv.ID_servicio
        FROM ejecucion_mantenimientos em
        LEFT JOIN sedes s ON s.id = em.sede_id
        LEFT JOIN servicios srv ON srv.id = em.servicio_id
        WHERE {where}
        ORDER BY em.fecha_ejecucion DESC, em.id DESC
        LIMIT ? OFFSET ?
        """,
        (*params, size_n, offset),
    ).fetchall()
    return {
        "ok": True,
        "total": int(total or 0),
        "page": page_n,
        "page_size": size_n,
        "items": [row_payload(r) for r in rows],
    }


def lookup_inventario_llave(conn, inventario_id: int):
    return conn.execute(
        """
        SELECT ie.id, ie.num_biomedica, ie.codigo_activo, ie.equipo, ie.freq_mp,
               ie.servicio_id, srv.sede_id, s.empresa_id, s.name_sede, srv.name_servicio
        FROM inventario_equipos ie
        JOIN servicios srv ON srv.id = ie.servicio_id
        JOIN sedes s ON s.id = srv.sede_id
        WHERE ie.id = ?
        """,
        (int(inventario_id),),
    ).fetchone()


def frecuencia_para_llave(
    conn,
    empresa_id: int,
    codigo_biomedica,
    codigo_activo,
    *,
    ge=None,
    pm_fabricante_meses=None,
    grupo_meses=None,
    frecuencia_tpm_label=None,
) -> dict:
    bio, act = llave_compuesta(codigo_biomedica, codigo_activo)
    stats = lookup_resumen_llave(resumen_ejecucion_por_llaves(conn, empresa_id), bio, act)
    return estructura_frecuencia_equipo(
        ge=ge,
        pm_fabricante_meses=pm_fabricante_meses,
        grupo_meses=grupo_meses,
        frecuencia_tpm_label=frecuencia_tpm_label,
        n_preventivos=stats.get("n_preventivos") or 0,
        n_correctivos=stats.get("n_correctivos") or 0,
        n_visitas=stats.get("n_visitas") or 0,
        primera_prev=stats.get("primera_prev"),
        ultima_prev=stats.get("ultima_prev"),
    )


def attach_pm_ejecucion(conn, equipos: list[dict], empresa_id: int, anio: int | None = None) -> None:
    try:
        resumen = resumen_ejecucion_por_llaves(conn, empresa_id, anio=anio)
    except Exception:
        resumen = {}
    attach_resumen_ejecucion(equipos, resumen)
    if anio is not None:
        try:
            from server.inventario import attach_flags_lista

            attach_flags_lista(equipos, resumen, anio)
        except Exception:
            pass


def _map_tipo_foscagib(raw, descripcion="") -> str | None:
    text = sanitize_string(str(raw or ""), 40).casefold()
    desc = sanitize_string(str(descripcion or ""), COMMENT).casefold()
    if "correct" in text:
        return "correctivo"
    if "prevent" in text:
        return "preventivo"
    if "correctiv" in desc:
        return "correctivo"
    if "preventiv" in desc or desc.startswith("mp "):
        return "preventivo"
    if text in {"", "otro", "calibración", "calibracion"}:
        return "visita"
    return _tipo(text) or "visita"


def _norm_nombre_equipo(value) -> str:
    text = sanitize_string(str(value or ""), LABEL).casefold()
    text = re.sub(r"[^a-z0-9áéíóúñü]+", " ", text)
    return " ".join(text.split())


def _unique_inventario(matches: list, empresa_id: int | None):
    pool = list(matches or [])
    if empresa_id:
        scoped = [row for row in pool if int(row["empresa_id"]) == int(empresa_id)]
        if scoped:
            pool = scoped
    if len(pool) == 1:
        return pool[0]
    return None


def _excel_date(value) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    parsed = parse_date(value)
    return parsed.isoformat() if parsed else None


def _norm_header(value) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip().casefold()


def _find_col(headers: list[str], *needles: str) -> int | None:
    for i, h in enumerate(headers):
        for needle in needles:
            if needle in h:
                return i
    return None


def _duracion_desde_horas(inicio, fin):
    """Compatibilidad: duración verificable = fin − inicio."""
    return duracion_desde_intervalo(inicio, fin)


def _pick_inventario(matches: list, empresa_id: int | None):
    if not matches:
        return None
    if empresa_id:
        for row in matches:
            if int(row["empresa_id"]) == int(empresa_id):
                return row
    return matches[0]


def import_foscagib_ejecucion(conn, xlsx_path: Path, *, empresa_id: int | None = None) -> dict:
    """Importa la hoja Data del Excel FOSCAGIB. No bloquea si el archivo no existe."""
    from openpyxl import load_workbook

    path = Path(xlsx_path)
    if not path.exists():
        return {"ok": False, "error": "archivo no encontrado", "inserted": 0}

    existing = {
        (r["codigo_biomedica_norm"], r["codigo_activo_norm"], r["tipo_mantenimiento"], r["fecha_ejecucion"])
        for r in conn.execute(
            """
            SELECT codigo_biomedica_norm, codigo_activo_norm, tipo_mantenimiento, fecha_ejecucion
            FROM ejecucion_mantenimientos
            WHERE fuente = 'importado'
            """
        ).fetchall()
    }

    inv_rows = conn.execute(
        """
        SELECT ie.num_biomedica, ie.codigo_activo, ie.equipo, ie.servicio_id, srv.sede_id, s.empresa_id
        FROM inventario_equipos ie
        JOIN servicios srv ON srv.id = ie.servicio_id
        JOIN sedes s ON s.id = srv.sede_id
        """
    ).fetchall()
    by_bio: dict[str, list] = {}
    by_act: dict[str, list] = {}
    by_nombre: dict[str, list] = {}
    for row in inv_rows:
        bio_key = normalize_codigo(row["num_biomedica"])
        act_key = normalize_codigo(row["codigo_activo"])
        nom_key = _norm_nombre_equipo(row["equipo"] if "equipo" in row.keys() else "")
        if bio_key:
            by_bio.setdefault(bio_key, []).append(row)
        if act_key:
            by_act.setdefault(act_key, []).append(row)
        if nom_key:
            by_nombre.setdefault(nom_key, []).append(row)

    default_empresa = int(empresa_id) if empresa_id else None
    if not default_empresa:
        row = conn.execute("SELECT empresa_id FROM sedes ORDER BY id LIMIT 1").fetchone()
        if row:
            default_empresa = int(row["empresa_id"])

    wb = load_workbook(path, read_only=True, data_only=True)
    if "Data" not in wb.sheetnames:
        return {"ok": False, "error": "hoja Data no encontrada", "inserted": 0}
    ws = wb["Data"]
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None) or ()
    headers = [_norm_header(x) for x in header_row]
    idx_atencion = _find_col(headers, "fecha de atenci", "fecha atenci")
    idx_cierre = _find_col(headers, "fecha de cierre", "fecha cierre")
    inserted = 0
    skipped = 0
    batch = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True)):
        if i >= IMPORT_ROWS * 2:
            break
        tipo_raw, fecha, descripcion, _actividad, _calidad, _ubicacion, responsable, cod_bio = row[:8]
        excel_norm = normalize_codigo(cod_bio)
        tipo = _map_tipo_foscagib(tipo_raw, descripcion)
        fecha_iso = _excel_date(row[11] if len(row) > 11 else None) or _excel_date(fecha)
        at_raw = row[idx_atencion] if idx_atencion is not None and idx_atencion < len(row) else None
        ci_raw = row[idx_cierre] if idx_cierre is not None and idx_cierre < len(row) else None
        duracion = duracion_desde_intervalo(at_raw, ci_raw)
        at_iso = _iso_datetime(parse_datetime(at_raw)) if duracion is not None else None
        ci_iso = _iso_datetime(parse_datetime(ci_raw)) if duracion is not None else None
        if not fecha_iso or not tipo:
            skipped += 1
            continue
        inv = None
        if excel_norm:
            inv = _pick_inventario(by_act.get(excel_norm) or [], empresa_id) or _pick_inventario(
                by_bio.get(excel_norm) or [], empresa_id
            )
        if not inv:
            nom = _norm_nombre_equipo(row[8] if len(row) > 8 else "")
            inv = _unique_inventario(by_nombre.get(nom) or [], empresa_id)
        if not excel_norm and not inv:
            skipped += 1
            continue
        if inv:
            bio_raw = (inv["num_biomedica"] or "").strip() if inv["num_biomedica"] not in (None, "") else ""
            act_raw = (inv["codigo_activo"] or "").strip() if inv["codigo_activo"] not in (None, "") else ""
            if not normalize_codigo(act_raw):
                act_raw = str(cod_bio).strip()
            if not normalize_codigo(bio_raw):
                bio_raw = ""
            emp = int(inv["empresa_id"])
            sede_id = inv["sede_id"]
            servicio_id = inv["servicio_id"]
        else:
            bio_raw = ""
            act_raw = str(cod_bio).strip()
            emp = int(empresa_id) if empresa_id else default_empresa
            sede_id = None
            servicio_id = None
        if not emp:
            skipped += 1
            continue
        bio_norm, act_norm = llave_compuesta(bio_raw, act_raw)
        if not bio_norm and not act_norm:
            skipped += 1
            continue
        dedup = (bio_norm, act_norm, tipo, fecha_iso)
        if dedup in existing:
            skipped += 1
            continue
        existing.add(dedup)
        obs = sanitize_string(str(descripcion or ""), COMMENT) or None
        proveedor = sanitize_string(str(responsable or ""), LABEL) or None
        batch.append(
            (
                bio_raw or bio_norm,
                act_raw or act_norm,
                bio_norm,
                act_norm,
                emp,
                sede_id,
                servicio_id,
                tipo,
                fecha_iso,
                at_iso,
                ci_iso,
                proveedor,
                duracion,
                obs,
                "importado",
            )
        )
        if len(batch) >= 400:
            _executemany_ejecucion(conn, batch)
            inserted += len(batch)
            batch.clear()
    if batch:
        _executemany_ejecucion(conn, batch)
        inserted += len(batch)
    return {"ok": True, "inserted": inserted, "skipped": skipped, "archivo": str(path)}


def _executemany_ejecucion(conn, batch: list[tuple]) -> None:
    conn.executemany(
        """
        INSERT INTO ejecucion_mantenimientos (
            codigo_biomedica, codigo_activo, codigo_biomedica_norm, codigo_activo_norm,
            empresa_id, sede_id, servicio_id, tipo_mantenimiento, fecha_ejecucion,
            fecha_atencion, fecha_cierre,
            proveedor, duracion_horas, observaciones, fuente, created_at, updated_at
        ) VALUES (
            ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, datetime('now'), datetime('now')
        )
        """,
        batch,
    )


def maybe_seed_foscagib_ejecucion(emp_conn) -> dict | None:
    return maybe_sync_foscagib_ejecucion(emp_conn)


def maybe_sync_foscagib_ejecucion(emp_conn) -> dict | None:
    """Importa FOSCAGIB si la tabla está vacía, o si SIGTB_SYNC_EJECUCION=1 (backfill)."""
    import os

    path = next((p for p in FOSCAGIB_EJECUCION_CANDIDATES if p.exists()), None)
    if not path:
        return None
    try:
        n = emp_conn.execute("SELECT COUNT(*) AS n FROM ejecucion_mantenimientos").fetchone()["n"]
    except Exception:
        return None
    force = os.environ.get("SIGTB_SYNC_EJECUCION", "").strip().lower() in {"1", "true", "yes"}
    if int(n or 0) > 0 and not force:
        return None
    try:
        summary = import_foscagib_ejecucion(emp_conn, path)
        print(
            f"[ejecucion] import FOSCAGIB: inserted={summary.get('inserted')} "
            f"skipped={summary.get('skipped')}"
        )
        return summary
    except Exception as exc:
        print(f"[ejecucion] import FOSCAGIB omitido: {exc}")
        return None
