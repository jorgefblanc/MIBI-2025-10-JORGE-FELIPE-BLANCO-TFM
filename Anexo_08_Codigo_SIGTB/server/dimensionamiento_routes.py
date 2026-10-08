"""API HTTP de dimensionamiento de personal de ingeniería clínica."""

from __future__ import annotations

import csv
import io
from datetime import datetime

from flask import Blueprint, jsonify, request, send_file, session
from openpyxl import Workbook

from server.pdf_staff import collect_org_staff

from server.authz import (
    assigned_empresa_id,
    can_access_sede,
    current_user,
    is_global_scope,
    login_required,
    permission_required,
    user_has_permission,
)
from server.db import get_empresas_connection, get_users_connection
from server.dimensionamiento import (
    CATALOGO_CRITICIDAD,
    CATALOGO_EQUIPOS,
    CATALOGO_SERVICIOS,
    CATALOGO_SI_NO,
    DEFAULT_EQUIPOS_POR_IC,
    DEFAULT_PARAMS,
    TOOLTIPS,
    _as_bool_si,
    _f,
    calc_fila_equipo,
    calc_ic_por_dotacion,
    calc_parametros,
    calc_resumen,
    criticidad_para_plantilla,
    calc_por_adquisicion,
    equipo_en_plantilla,
    normalize_criticidad,
    ratios_desde_params,
)
from server.adquisicion import (
    aplicar_mp_por_adquisicion,
    clasificar_forma,
    ensure_solicitudes_adquisicion,
    mp_forzado_por_adquisicion,
    _job_ve_tablero,
)
from server.inventario_params import (
    ACTIVITIES,
    TIP_MANUAL_SIN_HISTORICO,
    build_grupos_por_descripcion,
    descripcion_key,
    resolve_activity_for_calc,
)
from server.frecuencia_pm import (
    enrich_dim_fila_con_ejecucion,
    lookup_resumen_llave,
    resumen_ejecucion_por_llaves,
)
from server.inventario import apply_flags_equipo, clamp_anio
from server.parametros_equipo import get_by_descripcion, upsert_parametros
from server.limits import COMMENT
from server.formula_versions import formula_meta, registrar_calculo
from server.validators import sanitize_string

dim_bp = Blueprint("dimensionamiento", __name__, url_prefix="/api/dimensionamiento")


def _audit_dimensionamiento(
    *,
    empresa_id: int | None,
    sede_id: int | None = None,
    parametros: dict | None = None,
    resultado: dict | None = None,
) -> dict[str, str]:
    user = current_user()
    return registrar_calculo(
        modulo="dimensionamiento",
        empresa_id=empresa_id,
        sede_id=sede_id,
        usuario_id=(user or {}).get("id_usuario"),
        parametros=parametros,
        resultado=resultado,
        user=user,
    )


def _can_edit() -> bool:
    """Roles operativos y ADMIN editan; asistenciales solo consultan."""
    roll = (session.get("ROLL") or "").upper()
    if roll == "ADMIN":
        return True
    if roll == "OPERATIVO":
        return True
    if user_has_permission("admin_panel"):
        return True
    return False


def _can_edit_adicionales() -> bool:
    """Alta de actividades adicionales: solo cargos del ROLL OPERATIVO."""
    return (session.get("ROLL") or "").upper() == "OPERATIVO"


def _can_view_dim_empresa() -> bool:
    """Administrador, Director operativo y Coordinador operativo."""
    user = current_user()
    if not user:
        return False
    if is_global_scope(user) or (user.get("ROLL") or "").upper() == "ADMIN":
        return True
    if user_has_permission("admin_panel"):
        return True
    return _job_ve_tablero(user)


ADICIONAL_TIPOS = {
    "gestion_documental": "gestion_documental_h",
    "acompanamiento": "acompanamiento_h",
    "otras_tareas": "otras_tareas_h",
}


def _ensure_dim_ratio_schema(conn) -> None:
    """Columnas de plantilla por criticidad y unificación de 'Soporte vital' → Alto."""
    tables = {
        r["name"]
        for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table'").fetchall()
    }
    if "dim_parametros" in tables:
        cols = {
            c["name"] for c in conn.execute("PRAGMA table_info(dim_parametros)").fetchall()
        }
        additions = {
            "equipos_ic_bajo": "REAL NOT NULL DEFAULT 70",
            "equipos_ic_medio": "REAL NOT NULL DEFAULT 50",
            "equipos_ic_alto": "REAL NOT NULL DEFAULT 20",
        }
        for name, typ in additions.items():
            if name not in cols:
                conn.execute(f"ALTER TABLE dim_parametros ADD COLUMN {name} {typ}")
    if "dim_actividad_equipo" in tables:
        act_cols = {
            c["name"] for c in conn.execute("PRAGMA table_info(dim_actividad_equipo)").fetchall()
        }
        for name, typ in {
            "alto_unico_area": "INTEGER NOT NULL DEFAULT 0",
            "alto_especializado": "INTEGER NOT NULL DEFAULT 0",
        }.items():
            if name not in act_cols:
                conn.execute(f"ALTER TABLE dim_actividad_equipo ADD COLUMN {name} {typ}")
        conn.execute(
            """
            UPDATE dim_actividad_equipo
            SET criticidad = 'Alto'
            WHERE lower(trim(coalesce(criticidad, ''))) IN (
                    'soporte vital', 'soporte-vital', 'vital',
                    'especializado', 'especializados'
                )
               OR lower(coalesce(criticidad, '')) LIKE '%vital%'
               OR lower(coalesce(criticidad, '')) LIKE '%especializ%'
            """
        )


def _ensure_params(conn, empresa_id: int) -> dict:
    _ensure_dim_ratio_schema(conn)
    conn.execute(
        """
        INSERT OR IGNORE INTO dim_parametros (
            empresa_id, horas_semana, semanas_anio, vacaciones_h, festivos_h,
            permisos_h, productividad, correctivo_sin_historico, contingencia
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            empresa_id,
            DEFAULT_PARAMS["horas_semana"],
            DEFAULT_PARAMS["semanas_anio"],
            DEFAULT_PARAMS["vacaciones_h"],
            DEFAULT_PARAMS["festivos_h"],
            DEFAULT_PARAMS["permisos_h"],
            DEFAULT_PARAMS["productividad"],
            DEFAULT_PARAMS["correctivo_sin_historico"],
            DEFAULT_PARAMS["contingencia"],
        ),
    )
    row = conn.execute(
        "SELECT * FROM dim_parametros WHERE empresa_id = ?", (empresa_id,)
    ).fetchone()
    data = dict(row) if row else dict(DEFAULT_PARAMS)
    data.setdefault("equipos_ic_bajo", DEFAULT_EQUIPOS_POR_IC["Bajo"])
    data.setdefault("equipos_ic_medio", DEFAULT_EQUIPOS_POR_IC["Medio"])
    data.setdefault("equipos_ic_alto", DEFAULT_EQUIPOS_POR_IC["Alto"])
    if data.get("equipos_ic_bajo") in (None, ""):
        data["equipos_ic_bajo"] = DEFAULT_EQUIPOS_POR_IC["Bajo"]
    if data.get("equipos_ic_medio") in (None, ""):
        data["equipos_ic_medio"] = DEFAULT_EQUIPOS_POR_IC["Medio"]
    if data.get("equipos_ic_alto") in (None, ""):
        data["equipos_ic_alto"] = DEFAULT_EQUIPOS_POR_IC["Alto"]
    return data


def _sede_meta(conn, sede_id: int):
    return conn.execute(
        """
        SELECT s.id, s.ID_sede, s.name_sede, s.empresa_id, e.ID_Empresa, e.ID_NIT
        FROM sedes s
        JOIN empresas e ON e.id = s.empresa_id
        WHERE s.id = ?
        """,
        (sede_id,),
    ).fetchone()


def _empresa_accessible(empresa_id: int) -> bool:
    user = current_user()
    if is_global_scope(user) or (session.get("ROLL") or "").upper() == "ADMIN":
        return True
    return int(user.get("empresa_id") or 0) == int(empresa_id)


def _count_ingenieros_disponibles(sede_id: int, empresa_id: int | None = None) -> int:
    """Ingenieros/técnicos con cargo asignado a esa sede (usuario_sedes)."""
    with get_users_connection() as conn:
        rows = conn.execute(
            """
            SELECT DISTINCT u.id_usuario, u.ROLL, u.JOB
            FROM usuarios u
            JOIN usuario_sedes us ON us.usuario_id = u.id_usuario
            WHERE us.sede_id = ?
            """,
            (sede_id,),
        ).fetchall()
    return sum(1 for r in rows if _es_personal_ic(r["ROLL"], r["JOB"]))


def _es_personal_ic(roll, job) -> bool:
    roll = (roll or "").upper()
    job = (job or "").lower()
    return roll == "OPERATIVO" and (
        "ingenier" in job or "técnico" in job or "tecnico" in job
    )


def _count_ingenieros_empresa(empresa_id: int, sede_ids: list[int] | None = None) -> int:
    """Ingenieros/técnicos con cargo asignado a las sedes de la empresa."""
    if sede_ids is None:
        with get_empresas_connection() as emp_conn:
            sede_ids = [
                int(r["id"])
                for r in emp_conn.execute(
                    "SELECT id FROM sedes WHERE empresa_id = ?", (empresa_id,)
                ).fetchall()
            ]
    if not sede_ids:
        return 0
    seen: set[int] = set()
    with get_users_connection() as conn:
        placeholders = ",".join("?" * len(sede_ids))
        assigned = conn.execute(
            f"""
            SELECT DISTINCT u.id_usuario, u.ROLL, u.JOB
            FROM usuarios u
            JOIN usuario_sedes us ON us.usuario_id = u.id_usuario
            WHERE us.sede_id IN ({placeholders})
            """,
            sede_ids,
        ).fetchall()
        for r in assigned:
            if _es_personal_ic(r["ROLL"], r["JOB"]):
                seen.add(int(r["id_usuario"]))
    return len(seen)


def _list_empresas_scope(conn, user) -> list[dict]:
    if is_global_scope(user):
        rows = conn.execute(
            """
            SELECT id, ID_Empresa, ID_NIT
            FROM empresas
            ORDER BY ID_Empresa COLLATE NOCASE
            """
        ).fetchall()
    else:
        eid = assigned_empresa_id(user)
        if not eid:
            return []
        rows = conn.execute(
            """
            SELECT id, ID_Empresa, ID_NIT
            FROM empresas WHERE id = ?
            """,
            (eid,),
        ).fetchall()
    return [dict(r) for r in rows]


def _anio_dim() -> int:
    raw = request.args.get("anio")
    if raw in (None, "") and request.method in ("POST", "PUT", "PATCH"):
        body = request.get_json(silent=True) or {}
        raw = body.get("anio")
    return clamp_anio(raw)


def _empresa_tablero_payload(conn, empresa_id: int, anio: int | None = None) -> dict:
    empresa = conn.execute(
        "SELECT id, ID_Empresa, ID_NIT FROM empresas WHERE id = ?",
        (empresa_id,),
    ).fetchone()
    if not empresa:
        return {}
    raw_params = _ensure_params(conn, empresa_id)
    params_calc = calc_parametros(raw_params)
    ratios = ratios_desde_params(raw_params)
    sedes = conn.execute(
        """
        SELECT id, ID_sede, name_sede
        FROM sedes
        WHERE empresa_id = ?
        ORDER BY name_sede COLLATE NOCASE
        """,
        (empresa_id,),
    ).fetchall()

    all_equipos: list[dict] = []
    por_sede = []
    year = clamp_anio(anio)
    for sede in sedes:
        eqs = _load_equipos_contexto(conn, sede["id"], params_calc, anio=year)
        for eq in eqs:
            eq["sede_id"] = sede["id"]
            eq["ID_sede"] = sede["ID_sede"]
            eq["name_sede"] = sede["name_sede"]
        resumen_sede = calc_resumen(eqs, params_calc)
        ic_sede = calc_ic_por_dotacion(eqs, ratios)
        counts = {
            item["criticidad"]: item["equipos"]
            for item in ic_sede.get("por_criticidad_dotacion") or []
        }
        disponibles_sede = _count_ingenieros_disponibles(sede["id"], empresa_id)
        n_tercerizados = ic_sede.get("equipos_fuera_plantilla", 0)
        por_sede.append(
            {
                "sede_id": sede["id"],
                "ID_sede": sede["ID_sede"],
                "name_sede": sede["name_sede"],
                "total_equipos": resumen_sede["total_equipos"],
                "equipos_excluidos": resumen_sede["equipos_excluidos"],
                "equipos_plantilla": ic_sede["total_equipos_dotacion"],
                "equipos_comodato_leasing": n_tercerizados,
                "equipos_bajo": counts.get("Bajo", 0),
                "equipos_medio": counts.get("Medio", 0),
                "equipos_alto": counts.get("Alto", 0),
                "total_ajustado": resumen_sede["total_ajustado"],
                "ingenieros_requeridos": resumen_sede["ingenieros_requeridos"],
                "ingenieros_recomendados": resumen_sede["ingenieros_recomendados"],
                "ingenieros_por_dotacion": ic_sede["ingenieros_por_dotacion"],
                "ingenieros_recomendados_dotacion": ic_sede[
                    "ingenieros_recomendados_dotacion"
                ],
                "ingenieros_disponibles": disponibles_sede,
                "brecha_horas": disponibles_sede - resumen_sede["ingenieros_recomendados"],
                "brecha_dotacion": disponibles_sede
                - ic_sede["ingenieros_recomendados_dotacion"],
            }
        )
        all_equipos.extend(eqs)

    filas_empresa = []
    for eq in all_equipos:
        fila = dict(eq)
        sede_name = (eq.get("name_sede") or "").strip()
        srv = (eq.get("servicio") or "Sin servicio").strip() or "Sin servicio"
        fila["servicio"] = f"{sede_name} · {srv}" if sede_name else srv
        filas_empresa.append(fila)

    resumen = calc_resumen(filas_empresa, params_calc)
    ic = calc_ic_por_dotacion(all_equipos, ratios)
    por_adquisicion = calc_por_adquisicion(all_equipos)
    disponibles = _count_ingenieros_empresa(
        empresa_id, [int(s["id"]) for s in sedes]
    )
    return {
        "empresa": dict(empresa),
        "parametros": {**dict(raw_params), **params_calc},
        "ratios": ratios,
        "can_edit_ratios": _can_view_dim_empresa(),
        "resumen": resumen,
        "dotacion": ic,
        "por_adquisicion": por_adquisicion,
        "por_sede": por_sede,
        "anio": year,
        "dashboard": {
            **resumen,
            "total_equipos_dotacion": ic["total_equipos_dotacion"],
            "equipos_comodato_leasing": ic.get("equipos_fuera_plantilla", 0),
            "ingenieros_por_dotacion": ic["ingenieros_por_dotacion"],
            "ingenieros_recomendados_dotacion": ic["ingenieros_recomendados_dotacion"],
            "ingenieros_disponibles": disponibles,
            "brecha_horas": disponibles - resumen["ingenieros_recomendados"],
            "brecha_dotacion": disponibles - ic["ingenieros_recomendados_dotacion"],
            "sedes": len(sedes),
            "alto_unico_como_medio": ic.get("alto_unico_como_medio", 0),
            "alto_especializado": ic.get("alto_especializado", 0),
        },
        "tooltips": {
            "dashboard_empresa": TOOLTIPS["dashboard_empresa"],
            "equipos_ic_bajo": TOOLTIPS["equipos_ic_bajo"],
            "equipos_ic_medio": TOOLTIPS["equipos_ic_medio"],
            "equipos_ic_alto": TOOLTIPS["equipos_ic_alto"],
            "criticidad": TOOLTIPS["criticidad"],
            "ingenieros_requeridos": TOOLTIPS["ingenieros_requeridos"],
            "forma_adquisicion": TOOLTIPS["forma_adquisicion"],
        },
    }


def _actividad_defaults() -> dict:
    return {
        "criticidad": "Medio",
        "aplica_mp": 1,
        "freq_mp": 0,
        "tiempo_mp": 0,
        "aplica_cal": 0,
        "freq_cal": 0,
        "tiempo_cal": 0,
        "aplica_val": 0,
        "freq_val": 0,
        "tiempo_val": 0,
        "tercerizado_mp": 0,
        "tercerizado_cal": 0,
        "tercerizado_val": 0,
        "aplica_cap": 0,
        "freq_cap": 0,
        "tiempo_cap": 0,
        "tiene_historico_correctivo": 0,
        "historico_correctivo_h": None,
        "gestion_documental_h": 0,
        "acompanamiento_h": 0,
        "otras_tareas_h": 0,
        "observaciones": "",
        "alto_unico_area": 0,
        "alto_especializado": 0,
    }


def _map_criticidad(riesgo: str | None) -> str:
    raw = (riesgo or "").strip().lower()
    if not raw:
        return "Medio"
    if "soporte" in raw or "vital" in raw or "iib" in raw or "iii" in raw:
        return "Alto"
    if "alto" in raw or "alta" in raw or "iia" in raw:
        return "Alto"
    if "bajo" in raw or "baja" in raw or raw in ("i", "clase i"):
        return "Bajo"
    if "medio" in raw or "media" in raw:
        return "Medio"
    for c in CATALOGO_CRITICIDAD:
        if c.lower() == raw:
            return c
    return "Medio"


def _load_equipos_contexto(conn, sede_id: int, params_calc: dict, anio: int | None = None) -> list[dict]:
    """
    Carga equipos de la sede. Si un equipo no tiene MP/Cal/Val, reutiliza en memoria
    los datos de otro con la misma descripción (sin escribir en BD).
    """
    year = clamp_anio(anio)
    inv_rows = conn.execute(
        """
        SELECT ie.*, srv.name_servicio, srv.ID_servicio, srv.id AS servicio_id
        FROM inventario_equipos ie
        JOIN servicios srv ON srv.id = ie.servicio_id
        WHERE srv.sede_id = ?
        ORDER BY srv.name_servicio COLLATE NOCASE, ie.equipo COLLATE NOCASE, ie.id
        """,
        (sede_id,),
    ).fetchall()

    act_map = {
        r["inventario_equipo_id"]: dict(r)
        for r in conn.execute(
            """
            SELECT a.*
            FROM dim_actividad_equipo a
            JOIN inventario_equipos ie ON ie.id = a.inventario_equipo_id
            JOIN servicios srv ON srv.id = ie.servicio_id
            WHERE srv.sede_id = ?
            """,
            (sede_id,),
        ).fetchall()
    }

    inv_dicts = [dict(r) for r in inv_rows]
    correctivo_pct = _f(params_calc.get("correctivo_sin_historico"))
    empresa_row = conn.execute(
        "SELECT empresa_id FROM sedes WHERE id = ?",
        (sede_id,),
    ).fetchone()
    empresa_id = int(empresa_row["empresa_id"]) if empresa_row else 0
    try:
        resumen_ej = (
            resumen_ejecucion_por_llaves(conn, empresa_id, anio=year) if empresa_id else {}
        )
    except Exception:
        resumen_ej = {}
    result = []
    for inv in inv_dicts:
        act = act_map.get(inv["id"]) or _actividad_defaults()
        criticidad = normalize_criticidad(
            act.get("criticidad") or _map_criticidad(inv.get("clasificacion_riesgo"))
        )
        dim_fb = {
            **{f"aplica_{a}": act.get(f"aplica_{a}") for a in ACTIVITIES},
            **{f"freq_{a}": act.get(f"freq_{a}") for a in ACTIVITIES},
            **{f"tiempo_{a}": act.get(f"tiempo_{a}") for a in ACTIVITIES},
        }

        mp = resolve_activity_for_calc(inv, inv_dicts, "mp", dim_fallback=dim_fb)
        cal = resolve_activity_for_calc(inv, inv_dicts, "cal", dim_fallback=dim_fb)
        val = resolve_activity_for_calc(inv, inv_dicts, "val", dim_fallback=dim_fb)

        payload = {
            "inventario_equipo_id": inv["id"],
            "codigo": inv["num_biomedica"] or f"EQ-{inv['id']}",
            "servicio_id": inv["servicio_id"],
            "servicio": inv["name_servicio"],
            "equipo": inv["equipo"],
            "marca": inv["marca"],
            "modelo": inv["modelo"],
            "serial": inv["serie"],
            "ubicacion": inv["ubicacion"],
            "codigo_ubicacion": inv.get("codigo_ubicacion") or "",
            "ID_servicio": inv.get("ID_servicio") or "",
            "fecha_baja": inv.get("fecha_baja"),
            "fecha_ultimo_pm": inv.get("fecha_ultimo_pm"),
            "estado": inv["estado"],
            "criticidad": criticidad,
            "aplica_mp": mp["aplica"],
            "freq_mp": mp["freq"],
            "tiempo_mp": mp["tiempo"],
            "needs_manual_mp": mp["needs_manual"],
            "fuente_mp": mp["fuente"],
            "referencia_mp": mp.get("referencia_id"),
            "aplica_cal": cal["aplica"],
            "freq_cal": cal["freq"],
            "tiempo_cal": cal["tiempo"],
            "needs_manual_cal": cal["needs_manual"],
            "fuente_cal": cal["fuente"],
            "referencia_cal": cal.get("referencia_id"),
            "aplica_val": val["aplica"],
            "freq_val": val["freq"],
            "tiempo_val": val["tiempo"],
            "needs_manual_val": val["needs_manual"],
            "fuente_val": val["fuente"],
            "referencia_val": val.get("referencia_id"),
            "tercerizado_mp": bool(act.get("tercerizado_mp")),
            "tercerizado_cal": bool(act.get("tercerizado_cal")),
            "tercerizado_val": bool(act.get("tercerizado_val")),
            "aplica_cap": bool(act.get("aplica_cap")),
            "freq_cap": act.get("freq_cap") or 0,
            "tiempo_cap": act.get("tiempo_cap") or 0,
            "tiene_historico_correctivo": bool(act.get("tiene_historico_correctivo")),
            "historico_correctivo_h": act.get("historico_correctivo_h"),
            "gestion_documental_h": act.get("gestion_documental_h") or 0,
            "acompanamiento_h": act.get("acompanamiento_h") or 0,
            "otras_tareas_h": act.get("otras_tareas_h") or 0,
            "observaciones": act.get("observaciones") or "",
            "alto_unico_area": bool(act.get("alto_unico_area")),
            "alto_especializado": bool(act.get("alto_especializado")),
            "forma_adquisicion": inv.get("forma_adquisicion") or "",
            "adquisicion_clase": clasificar_forma(inv.get("forma_adquisicion")),
            "mp_forzado_tercerizado": mp_forzado_por_adquisicion(inv.get("forma_adquisicion")),
            "num_biomedica": inv.get("num_biomedica") or "",
            "codigo_activo": inv.get("codigo_activo") or "",
        }
        aplicar_mp_por_adquisicion(payload)
        stats_ej = lookup_resumen_llave(
            resumen_ej,
            payload.get("num_biomedica"),
            payload.get("codigo_activo"),
        )
        apply_flags_equipo(payload, stats_ej, year)
        enrich_dim_fila_con_ejecucion(payload, stats_ej)
        if payload["criticidad"] == "Alto" and not payload["alto_unico_area"] and not payload["alto_especializado"]:
            mapped = _map_criticidad(inv.get("clasificacion_riesgo"))
            if mapped == "Alto":
                payload["alto_especializado"] = True
        calc = calc_fila_equipo(payload, correctivo_pct, anio=year)
        payload.update(calc)
        payload["en_dotacion"] = calc.get("en_dotacion", True)
        payload["en_plantilla"] = calc.get("en_plantilla", equipo_en_plantilla(payload))
        payload["criticidad_plantilla"] = criticidad_para_plantilla(payload)
        result.append(payload)
    return result


@dim_bp.get("/meta")
@login_required
@permission_required("view_dimensionamiento")
def meta():
    return jsonify(
        {
            "ok": True,
            "tooltips": {**TOOLTIPS, "manual_sin_historico": TIP_MANUAL_SIN_HISTORICO},
            "catalogos": {
                "si_no": CATALOGO_SI_NO,
                "criticidad": CATALOGO_CRITICIDAD,
                "servicios": CATALOGO_SERVICIOS,
                "equipos": CATALOGO_EQUIPOS,
            },
            "can_edit": _can_edit(),
            "can_edit_adicionales": _can_edit_adicionales(),
            "perfil": DEFAULT_PARAMS["perfil"],
            "formula": formula_meta("dimensionamiento"),
        }
    )


@dim_bp.get("/parametros")
@login_required
@permission_required("view_dimensionamiento")
def get_parametros():
    empresa_id = request.args.get("empresa_id", type=int)
    if not empresa_id:
        return jsonify({"ok": False, "error": "empresa_id es obligatorio."}), 400
    if not _empresa_accessible(empresa_id):
        return jsonify({"ok": False, "error": "No autorizado para esta empresa."}), 403

    with get_empresas_connection() as conn:
        raw = _ensure_params(conn, empresa_id)
        conn.commit()
    calc = calc_parametros(raw)
    return jsonify({"ok": True, "parametros": {**raw, **calc}, "can_edit": _can_edit()})


@dim_bp.put("/parametros")
@login_required
@permission_required("view_dimensionamiento")
def put_parametros():
    if not _can_edit():
        return jsonify(
            {
                "ok": False,
                "error": "Solo roles operativos pueden editar parámetros de dimensionamiento.",
            }
        ), 403

    body = request.get_json(silent=True) or {}
    empresa_id = body.get("empresa_id")
    try:
        empresa_id = int(empresa_id)
    except (TypeError, ValueError):
        return jsonify({"ok": False, "error": "empresa_id inválido."}), 400
    if not _empresa_accessible(empresa_id):
        return jsonify({"ok": False, "error": "No autorizado para esta empresa."}), 403

    fields = {
        "horas_semana": _f(body.get("horas_semana"), DEFAULT_PARAMS["horas_semana"]),
        "semanas_anio": _f(body.get("semanas_anio"), DEFAULT_PARAMS["semanas_anio"]),
        "vacaciones_h": _f(body.get("vacaciones_h"), DEFAULT_PARAMS["vacaciones_h"]),
        "festivos_h": _f(body.get("festivos_h"), DEFAULT_PARAMS["festivos_h"]),
        "permisos_h": _f(body.get("permisos_h"), DEFAULT_PARAMS["permisos_h"]),
        "productividad": _f(body.get("productividad"), DEFAULT_PARAMS["productividad"]),
        "correctivo_sin_historico": _f(
            body.get("correctivo_sin_historico"),
            DEFAULT_PARAMS["correctivo_sin_historico"],
        ),
        "contingencia": _f(body.get("contingencia"), DEFAULT_PARAMS["contingencia"]),
    }
    if fields["productividad"] < 0 or fields["productividad"] > 1:
        return jsonify({"ok": False, "error": "productividad debe estar entre 0 y 1."}), 400
    if fields["correctivo_sin_historico"] < 0 or fields["contingencia"] < 0:
        return jsonify({"ok": False, "error": "Porcentajes no pueden ser negativos."}), 400

    with get_empresas_connection() as conn:
        _ensure_params(conn, empresa_id)
        conn.execute(
            """
            UPDATE dim_parametros SET
                horas_semana = ?, semanas_anio = ?, vacaciones_h = ?, festivos_h = ?,
                permisos_h = ?, productividad = ?, correctivo_sin_historico = ?,
                contingencia = ?, updated_at = datetime('now')
            WHERE empresa_id = ?
            """,
            (
                fields["horas_semana"],
                fields["semanas_anio"],
                fields["vacaciones_h"],
                fields["festivos_h"],
                fields["permisos_h"],
                fields["productividad"],
                fields["correctivo_sin_historico"],
                fields["contingencia"],
                empresa_id,
            ),
        )
        conn.commit()
        raw = _ensure_params(conn, empresa_id)

    calc = calc_parametros(raw)
    meta = _audit_dimensionamiento(
        empresa_id=empresa_id,
        parametros={"accion": "put_parametros", **fields},
        resultado={
            "h_nom": calc.get("horas_nominales"),
            "h_disp": calc.get("horas_efectivas"),
        },
    )
    return jsonify(
        {
            "ok": True,
            "message": "Parámetros actualizados.",
            "parametros": {**raw, **calc},
            "formula": meta,
        }
    )


@dim_bp.get("/empresa/tablero")
@login_required
@permission_required("view_dimensionamiento_empresa", "admin_panel")
def get_empresa_tablero():
    if not _can_view_dim_empresa():
        return jsonify(
            {
                "ok": False,
                "error": "Este tablero está reservado a Administrador, Coordinador y Director operativo.",
            }
        ), 403

    user = current_user()
    empresa_id = request.args.get("empresa_id", type=int)
    if not is_global_scope(user):
        assigned = assigned_empresa_id(user)
        if not assigned:
            return jsonify(
                {"ok": False, "error": "Tu usuario aún no tiene empresa asignada."}
            ), 400
        if empresa_id and int(empresa_id) != int(assigned):
            return jsonify({"ok": False, "error": "No autorizado para esta empresa."}), 403
        empresa_id = int(assigned)
    elif empresa_id and not _empresa_accessible(empresa_id):
        return jsonify({"ok": False, "error": "No autorizado para esta empresa."}), 403

    with get_empresas_connection() as conn:
        empresas = _list_empresas_scope(conn, user)
        if not empresa_id:
            if len(empresas) == 1:
                empresa_id = int(empresas[0]["id"])
            else:
                return jsonify(
                    {
                        "ok": True,
                        "need_empresa": True,
                        "empresas": empresas,
                        "message": "Selecciona una empresa para consolidar el dimensionamiento.",
                    }
                )
        payload = _empresa_tablero_payload(conn, int(empresa_id), anio=_anio_dim())
        conn.commit()

    if not payload:
        return jsonify({"ok": False, "error": "Empresa no encontrada."}), 404
    return jsonify({"ok": True, "need_empresa": False, "empresas": empresas, **payload})


@dim_bp.put("/empresa/ratios")
@login_required
@permission_required("view_dimensionamiento_empresa", "admin_panel")
def put_empresa_ratios():
    if not _can_view_dim_empresa():
        return jsonify(
            {
                "ok": False,
                "error": "Solo Administrador, Coordinador y Director operativo pueden editar estos parámetros.",
            }
        ), 403

    body = request.get_json(silent=True) or {}
    user = current_user()
    try:
        empresa_id = int(body.get("empresa_id") or 0)
    except (TypeError, ValueError):
        empresa_id = 0
    if not is_global_scope(user):
        assigned = assigned_empresa_id(user)
        if not assigned:
            return jsonify(
                {"ok": False, "error": "Tu usuario aún no tiene empresa asignada."}
            ), 400
        empresa_id = int(assigned)
    elif not empresa_id:
        return jsonify({"ok": False, "error": "empresa_id es obligatorio."}), 400
    if not _empresa_accessible(empresa_id):
        return jsonify({"ok": False, "error": "No autorizado para esta empresa."}), 403

    bajo = _f(body.get("equipos_ic_bajo"), DEFAULT_EQUIPOS_POR_IC["Bajo"])
    medio = _f(body.get("equipos_ic_medio"), DEFAULT_EQUIPOS_POR_IC["Medio"])
    alto = _f(body.get("equipos_ic_alto"), DEFAULT_EQUIPOS_POR_IC["Alto"])
    if min(bajo, medio, alto) < 1:
        return jsonify(
            {
                "ok": False,
                "error": "Cada ratio debe ser al menos 1 equipo por ingeniero clínico.",
            }
        ), 400

    with get_empresas_connection() as conn:
        _ensure_params(conn, empresa_id)
        conn.execute(
            """
            UPDATE dim_parametros SET
                equipos_ic_bajo = ?, equipos_ic_medio = ?, equipos_ic_alto = ?,
                updated_at = datetime('now')
            WHERE empresa_id = ?
            """,
            (bajo, medio, alto, empresa_id),
        )
        conn.commit()
        raw = _ensure_params(conn, empresa_id)

    meta = _audit_dimensionamiento(
        empresa_id=empresa_id,
        parametros={
            "accion": "put_empresa_ratios",
            "equipos_ic_bajo": bajo,
            "equipos_ic_medio": medio,
            "equipos_ic_alto": alto,
        },
        resultado=ratios_desde_params(raw),
    )
    return jsonify(
        {
            "ok": True,
            "message": "Parámetros de plantilla actualizados.",
            "ratios": ratios_desde_params(raw),
            "parametros": raw,
            "formula": meta,
        }
    )


def _load_tablero_empresa_actual():
    """Carga el consolidado empresarial o una respuesta de error Flask."""
    if not _can_view_dim_empresa():
        return jsonify(
            {
                "ok": False,
                "error": "Este tablero está reservado a Administrador, Coordinador y Director operativo.",
            }
        ), 403
    user = current_user()
    empresa_id = request.args.get("empresa_id", type=int)
    if not is_global_scope(user):
        assigned = assigned_empresa_id(user)
        if not assigned:
            return jsonify(
                {"ok": False, "error": "Tu usuario aún no tiene empresa asignada."}
            ), 400
        if empresa_id and int(empresa_id) != int(assigned):
            return jsonify({"ok": False, "error": "No autorizado para esta empresa."}), 403
        empresa_id = int(assigned)
    elif empresa_id and not _empresa_accessible(empresa_id):
        return jsonify({"ok": False, "error": "No autorizado para esta empresa."}), 403
    with get_empresas_connection() as conn:
        empresas = _list_empresas_scope(conn, user)
        if not empresa_id:
            if len(empresas) == 1:
                empresa_id = int(empresas[0]["id"])
            else:
                return jsonify(
                    {
                        "ok": False,
                        "error": "Selecciona una empresa para generar el informe.",
                        "need_empresa": True,
                        "empresas": empresas,
                    }
                ), 400
        payload = _empresa_tablero_payload(conn, int(empresa_id), anio=_anio_dim())
        conn.commit()
    if not payload:
        return jsonify({"ok": False, "error": "Empresa no encontrada."}), 404
    payload["empresas"] = empresas
    return payload


@dim_bp.get("/empresa/export.pdf")
@login_required
@permission_required("view_dimensionamiento_empresa", "admin_panel")
def export_empresa_pdf():
    loaded = _load_tablero_empresa_actual()
    if not isinstance(loaded, dict):
        return loaded
    from server.dimensionamiento_pdf import build_informe_empresa_pdf

    emp = loaded.get("empresa") or {}
    loaded["staff"] = collect_org_staff(empresa_id=emp.get("id"))
    pdf_bytes = build_informe_empresa_pdf(loaded)
    data = io.BytesIO(pdf_bytes)
    data.seek(0)
    emp = loaded.get("empresa") or {}
    fname = f"informe_dimensionamiento_{emp.get('ID_Empresa') or emp.get('id')}_{datetime.now():%Y%m%d_%H%M%S}.pdf"
    return send_file(data, mimetype="application/pdf", as_attachment=True, download_name=fname)


@dim_bp.get("/empresa/export.xlsx")
@login_required
@permission_required("view_dimensionamiento_empresa", "admin_panel")
def export_empresa_xlsx():
    loaded = _load_tablero_empresa_actual()
    if not isinstance(loaded, dict):
        return loaded
    data = _build_empresa_xlsx(loaded)
    emp = loaded.get("empresa") or {}
    fname = f"dimensionamiento_{emp.get('ID_Empresa') or emp.get('id')}_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    return send_file(
        data,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=fname,
    )


def _build_empresa_xlsx(payload: dict) -> io.BytesIO:
    dash = payload.get("dashboard") or {}
    wb = Workbook()
    ws = wb.active
    ws.title = "Dashboard"
    ws.append(["Indicador", "Valor"])
    ws.append(["Empresa", (payload.get("empresa") or {}).get("ID_Empresa")])
    ws.append(["Sedes", dash.get("sedes")])
    ws.append(["Equipos en carga horaria", dash.get("total_equipos")])
    ws.append(["Equipos en plantilla (criticidad)", dash.get("total_equipos_dotacion")])
    ws.append(["Comodato/Leasing (fuera de plantilla)", dash.get("equipos_comodato_leasing")])
    ws.append(["Horas ajustadas / año", dash.get("total_ajustado")])
    ws.append(["IC por carga horaria (decimal)", dash.get("ingenieros_requeridos")])
    ws.append(["IC por carga horaria (recomendados)", dash.get("ingenieros_recomendados")])
    ws.append(["IC por criticidad (decimal)", dash.get("ingenieros_por_dotacion")])
    ws.append(["IC por criticidad (recomendados)", dash.get("ingenieros_recomendados_dotacion")])
    ws.append(["IC/técnicos asignados", dash.get("ingenieros_disponibles")])
    ws.append(["Altos únicos con ratio medio", dash.get("alto_unico_como_medio")])

    ws_adq = wb.create_sheet("Adquisicion")
    ws_adq.append(["Forma de adquisición", "Equipos", "En plantilla", "Fuera de plantilla", "Horas carga"])
    for row in payload.get("por_adquisicion") or []:
        ws_adq.append(
            [
                row.get("forma_adquisicion"),
                row.get("equipos"),
                row.get("en_plantilla"),
                row.get("fuera_plantilla"),
                row.get("horas"),
            ]
        )

    ws_c = wb.create_sheet("Criticidad")
    ws_c.append(["Criticidad", "Equipos", "Equipos / IC", "IC decimal", "IC recomendados"])
    for row in (payload.get("dotacion") or {}).get("por_criticidad_dotacion") or []:
        ws_c.append(
            [
                row.get("criticidad"),
                row.get("equipos"),
                row.get("equipos_por_ic"),
                row.get("ingenieros"),
                row.get("ingenieros_recomendados"),
            ]
        )

    ws_a = wb.create_sheet("Actividades")
    ws_a.append(["Actividad", "Horas/año"])
    for row in (payload.get("resumen") or {}).get("por_actividad") or []:
        ws_a.append([row.get("actividad"), row.get("horas")])

    ws_s = wb.create_sheet("Sedes")
    ws_s.append(
        [
            "Sede",
            "Equipos carga",
            "Plantilla",
            "Comodato/Leasing",
            "Bajo",
            "Medio",
            "Alto",
            "Horas ajustadas",
            "IC horas",
            "IC criticidad",
            "Asignados",
        ]
    )
    for row in payload.get("por_sede") or []:
        ws_s.append(
            [
                f"{row.get('ID_sede') or ''} {row.get('name_sede') or ''}".strip(),
                row.get("total_equipos"),
                row.get("equipos_plantilla"),
                row.get("equipos_comodato_leasing"),
                row.get("equipos_bajo"),
                row.get("equipos_medio"),
                row.get("equipos_alto"),
                row.get("total_ajustado"),
                row.get("ingenieros_recomendados"),
                row.get("ingenieros_recomendados_dotacion"),
                row.get("ingenieros_disponibles"),
            ]
        )

    ws_t = wb.create_sheet("Tipos")
    ws_t.append(["Equipo", "Cantidad", "Criticidad", "Equipos / IC", "IC decimal", "IC recomendados"])
    for row in (payload.get("dotacion") or {}).get("por_tipo_equipo") or []:
        ws_t.append(
            [
                row.get("equipo"),
                row.get("equipos"),
                row.get("criticidad"),
                row.get("equipos_por_ic"),
                row.get("ingenieros"),
                row.get("ingenieros_recomendados"),
            ]
        )

    data = io.BytesIO()
    wb.save(data)
    data.seek(0)
    return data


@dim_bp.get("/contexto")
@login_required
@permission_required("view_dimensionamiento")
def get_contexto():
    sede_id = request.args.get("sede_id", type=int)
    if not sede_id:
        return jsonify({"ok": False, "error": "sede_id es obligatorio."}), 400

    user = current_user()
    if not can_access_sede(user, sede_id):
        return jsonify({"ok": False, "error": "No autorizado para esta sede."}), 403

    with get_empresas_connection() as conn:
        sede = _sede_meta(conn, sede_id)
        if not sede:
            return jsonify({"ok": False, "error": "Sede no encontrada."}), 404
        raw_params = _ensure_params(conn, sede["empresa_id"])
        params_calc = calc_parametros(raw_params)
        year = _anio_dim()
        equipos = _load_equipos_contexto(conn, sede_id, params_calc, anio=year)
        grupos = build_grupos_por_descripcion(equipos)
        for g in grupos:
            saved = get_by_descripcion(g.get("equipo") or g.get("grupo_key") or "")
            g["catalogo_disponible"] = bool(saved)
            g["catalogo_parametros"] = saved

    try:
        adquisicion_inbox = ensure_solicitudes_adquisicion(
            user=user,
            sede=dict(sede),
            equipos=equipos,
        )
    except Exception as exc:
        adquisicion_inbox = {
            "created": 0,
            "already": 0,
            "faltantes": 0,
            "warning": f"No se pudieron radicar solicitudes de adquisición: {exc}",
        }

    faltantes = int(adquisicion_inbox.get("faltantes") or 0)
    resumen = calc_resumen(equipos, params_calc)
    resumen["definitivo"] = faltantes == 0
    resumen["motivo_no_definitivo"] = (
        None
        if faltantes == 0
        else (
            f"Hay {faltantes} equipo(s) sin forma de adquisición. "
            "El dimensionamiento se muestra como preliminar hasta completar el dato."
        )
    )
    disponibles = _count_ingenieros_disponibles(sede_id, sede["empresa_id"])
    brecha = disponibles - resumen["ingenieros_recomendados"]

    return jsonify(
        {
            "ok": True,
            "can_edit": _can_edit(),
            "can_edit_adicionales": _can_edit_adicionales(),
            "anio": year,
            "sede": {
                "id": sede["id"],
                "ID_sede": sede["ID_sede"],
                "name_sede": sede["name_sede"],
                "empresa_id": sede["empresa_id"],
                "ID_Empresa": sede["ID_Empresa"],
                "ID_NIT": sede["ID_NIT"],
            },
            "parametros": {**dict(raw_params), **params_calc},
            "equipos": equipos,
            "grupos": grupos,
            "adquisicion_inbox": adquisicion_inbox,
            "resumen": resumen,
            "dashboard": {
                **resumen,
                "ingenieros_disponibles": disponibles,
                "brecha_personal": brecha,
            },
            "informe": {
                "objetivo": TOOLTIPS["objetivo"],
                "perfil": params_calc.get("perfil"),
                "total_equipos": resumen["total_equipos"],
                "total_base": resumen["total_base"],
                "contingencia_h": resumen["contingencia_h"],
                "total_ajustado": resumen["total_ajustado"],
                "horas_efectivas": resumen["horas_efectivas"],
                "ingenieros_requeridos": resumen["ingenieros_requeridos"],
                "ingenieros_recomendados": resumen["ingenieros_recomendados"],
                "ingenieros_disponibles": disponibles,
                "por_servicio": resumen["por_servicio"],
                "interpretacion": (
                    "Si la carga total ajustada supera la capacidad efectiva del equipo actual, "
                    "se genera riesgo de incumplimiento del mantenimiento programado, retrasos "
                    "en correctivos, debilidad documental y menor capacidad para auditorías."
                ),
                "recomendacion": (
                    "Usar este resultado como soporte técnico para justificar contratación, "
                    "redistribución de funciones, tercerización controlada o fortalecimiento "
                    "del área de Ingeniería Clínica."
                ),
            },
            "formula": formula_meta("dimensionamiento"),
        }
    )


@dim_bp.get("/parametros-tipo")
@login_required
@permission_required("view_dimensionamiento")
def get_parametros_tipo():
    descripcion = sanitize_string(request.args.get("descripcion")) or ""
    if not descripcion:
        return jsonify({"ok": False, "error": "descripcion es obligatoria."}), 400
    saved = get_by_descripcion(descripcion)
    return jsonify(
        {
            "ok": True,
            "found": bool(saved),
            "parametros": saved,
        }
    )


@dim_bp.put("/parametros-tipo")
@login_required
@permission_required("view_dimensionamiento")
def put_parametros_tipo():
    if not _can_edit():
        return jsonify(
            {"ok": False, "error": "Solo roles operativos pueden guardar el catálogo."}
        ), 403
    body = request.get_json(silent=True) or {}
    descripcion = sanitize_string(body.get("descripcion")) or ""
    if not descripcion:
        return jsonify({"ok": False, "error": "descripcion es obligatoria."}), 400
    try:
        saved = upsert_parametros(descripcion, body)
    except ValueError as err:
        return jsonify({"ok": False, "error": str(err)}), 400
    return jsonify(
        {
            "ok": True,
            "message": "Parámetros del tipo de equipo actualizados.",
            "parametros": saved,
        }
    )


def _parse_actividad_body(body: dict) -> dict:
    criticidad = normalize_criticidad(sanitize_string(body.get("criticidad")) or "Medio")
    alto_unico = 1 if _as_bool_si(body.get("alto_unico_area")) else 0
    alto_esp = 1 if _as_bool_si(body.get("alto_especializado")) else 0
    if criticidad != "Alto":
        alto_unico = 0
        alto_esp = 0
    data = {
        "criticidad": criticidad,
        "alto_unico_area": alto_unico,
        "alto_especializado": alto_esp,
        "aplica_mp": 1 if _as_bool_si(body.get("aplica_mp")) else 0,
        "freq_mp": _f(body.get("freq_mp")),
        "tiempo_mp": _f(body.get("tiempo_mp")),
        "aplica_cal": 1 if _as_bool_si(body.get("aplica_cal")) else 0,
        "freq_cal": _f(body.get("freq_cal")),
        "tiempo_cal": _f(body.get("tiempo_cal")),
        "aplica_val": 1 if _as_bool_si(body.get("aplica_val")) else 0,
        "freq_val": _f(body.get("freq_val")),
        "tiempo_val": _f(body.get("tiempo_val")),
        "tercerizado_mp": 1 if _as_bool_si(body.get("tercerizado_mp")) else 0,
        "tercerizado_cal": 1 if _as_bool_si(body.get("tercerizado_cal")) else 0,
        "tercerizado_val": 1 if _as_bool_si(body.get("tercerizado_val")) else 0,
        "aplica_cap": 1 if _as_bool_si(body.get("aplica_cap")) else 0,
        "freq_cap": _f(body.get("freq_cap")),
        "tiempo_cap": _f(body.get("tiempo_cap")),
        "tiene_historico_correctivo": 1
        if _as_bool_si(body.get("tiene_historico_correctivo"))
        else 0,
        "historico_correctivo_h": (
            None
            if body.get("historico_correctivo_h") in (None, "")
            else _f(body.get("historico_correctivo_h"))
        ),
        "gestion_documental_h": _f(body.get("gestion_documental_h")),
        "acompanamiento_h": _f(body.get("acompanamiento_h")),
        "otras_tareas_h": _f(body.get("otras_tareas_h")),
        "observaciones": sanitize_string(body.get("observaciones"), COMMENT) or "",
    }
    _zero_freq_tiempo_tercerizado(data)
    return data


def _zero_freq_tiempo_tercerizado(data: dict) -> None:
    """Servicio tercerizado: Freq/año y Tiempo (h) quedan en 0 y no alimentan carga IC."""
    for prefix in ("mp", "cal", "val"):
        if data.get(f"tercerizado_{prefix}"):
            data[f"freq_{prefix}"] = 0.0
            data[f"tiempo_{prefix}"] = 0.0


@dim_bp.put("/equipos/<int:inventario_equipo_id>")
@login_required
@permission_required("view_dimensionamiento")
def put_equipo_actividad(inventario_equipo_id: int):
    if not _can_edit():
        return jsonify(
            {"ok": False, "error": "Solo roles operativos pueden editar actividades."}
        ), 403

    body = request.get_json(silent=True) or {}
    data = _parse_actividad_body(body)
    if data["criticidad"] not in CATALOGO_CRITICIDAD:
        return jsonify({"ok": False, "error": "Criticidad inválida."}), 400
    if data["criticidad"] == "Alto" and not data["alto_unico_area"] and not data["alto_especializado"]:
        return jsonify(
            {
                "ok": False,
                "error": (
                    "Si la criticidad es Alta, indique si es por único equipo en el área, "
                    "por alta complejidad / especializado, o ambas."
                ),
            }
        ), 400

    user = current_user()
    with get_empresas_connection() as conn:
        _ensure_dim_ratio_schema(conn)
        inv = conn.execute(
            """
            SELECT ie.id, ie.equipo, ie.forma_adquisicion, ie.estado, ie.fecha_baja,
                   ie.ubicacion, ie.codigo_ubicacion, ie.fecha_ultimo_pm,
                   srv.sede_id, srv.ID_servicio, s.empresa_id
            FROM inventario_equipos ie
            JOIN servicios srv ON srv.id = ie.servicio_id
            JOIN sedes s ON s.id = srv.sede_id
            WHERE ie.id = ?
            """,
            (inventario_equipo_id,),
        ).fetchone()
        if not inv:
            return jsonify({"ok": False, "error": "Equipo no encontrado en inventario."}), 404
        if not can_access_sede(user, inv["sede_id"]):
            return jsonify({"ok": False, "error": "No autorizado."}), 403

        if mp_forzado_por_adquisicion(inv["forma_adquisicion"]):
            data["aplica_mp"] = 1
            data["tercerizado_mp"] = 1
            _zero_freq_tiempo_tercerizado(data)

        conn.execute(
            """
            INSERT INTO dim_actividad_equipo (
                inventario_equipo_id, criticidad, alto_unico_area, alto_especializado,
                aplica_mp, freq_mp, tiempo_mp,
                aplica_cal, freq_cal, tiempo_cal,
                aplica_val, freq_val, tiempo_val,
                tercerizado_mp, tercerizado_cal, tercerizado_val,
                aplica_cap, freq_cap, tiempo_cap,
                tiene_historico_correctivo, historico_correctivo_h,
                gestion_documental_h, acompanamiento_h, otras_tareas_h,
                observaciones, updated_at
            ) VALUES (
                ?, ?, ?, ?, ?, ?, ?,
                ?, ?, ?, ?, ?, ?,
                ?, ?, ?, ?, ?, ?,
                ?, ?, ?, ?, ?, ?,
                datetime('now')
            )
            ON CONFLICT(inventario_equipo_id) DO UPDATE SET
                criticidad = excluded.criticidad,
                alto_unico_area = excluded.alto_unico_area,
                alto_especializado = excluded.alto_especializado,
                aplica_mp = excluded.aplica_mp,
                freq_mp = excluded.freq_mp,
                tiempo_mp = excluded.tiempo_mp,
                aplica_cal = excluded.aplica_cal,
                freq_cal = excluded.freq_cal,
                tiempo_cal = excluded.tiempo_cal,
                aplica_val = excluded.aplica_val,
                freq_val = excluded.freq_val,
                tiempo_val = excluded.tiempo_val,
                tercerizado_mp = excluded.tercerizado_mp,
                tercerizado_cal = excluded.tercerizado_cal,
                tercerizado_val = excluded.tercerizado_val,
                aplica_cap = excluded.aplica_cap,
                freq_cap = excluded.freq_cap,
                tiempo_cap = excluded.tiempo_cap,
                tiene_historico_correctivo = excluded.tiene_historico_correctivo,
                historico_correctivo_h = excluded.historico_correctivo_h,
                gestion_documental_h = excluded.gestion_documental_h,
                acompanamiento_h = excluded.acompanamiento_h,
                otras_tareas_h = excluded.otras_tareas_h,
                observaciones = excluded.observaciones,
                updated_at = datetime('now')
            """,
            (
                inventario_equipo_id,
                data["criticidad"],
                data["alto_unico_area"],
                data["alto_especializado"],
                data["aplica_mp"],
                data["freq_mp"],
                data["tiempo_mp"],
                data["aplica_cal"],
                data["freq_cal"],
                data["tiempo_cal"],
                data["aplica_val"],
                data["freq_val"],
                data["tiempo_val"],
                data["tercerizado_mp"],
                data["tercerizado_cal"],
                data["tercerizado_val"],
                data["aplica_cap"],
                data["freq_cap"],
                data["tiempo_cap"],
                data["tiene_historico_correctivo"],
                data["historico_correctivo_h"],
                data["gestion_documental_h"],
                data["acompanamiento_h"],
                data["otras_tareas_h"],
                data["observaciones"],
            ),
        )

        peers = conn.execute(
            """
            SELECT ie.id
            FROM inventario_equipos ie
            JOIN servicios srv ON srv.id = ie.servicio_id
            WHERE srv.sede_id = ?
              AND ie.equipo = ? COLLATE NOCASE
              AND ie.id != ?
            """,
            (inv["sede_id"], inv["equipo"], inventario_equipo_id),
        ).fetchall()
        for peer in peers:
            conn.execute(
                """
                INSERT INTO dim_actividad_equipo (
                    inventario_equipo_id, criticidad, alto_unico_area, alto_especializado
                ) VALUES (?, ?, ?, ?)
                ON CONFLICT(inventario_equipo_id) DO UPDATE SET
                    criticidad = excluded.criticidad,
                    alto_unico_area = excluded.alto_unico_area,
                    alto_especializado = excluded.alto_especializado,
                    updated_at = datetime('now')
                """,
                (
                    peer["id"],
                    data["criticidad"],
                    data["alto_unico_area"],
                    data["alto_especializado"],
                ),
            )

        # Guarda MP/Cal/Val solo en este equipo (no propaga a similares).
        conn.execute(
            """
            UPDATE inventario_equipos SET
                aplica_mp = ?, freq_mp = ?, tiempo_mp = ?,
                aplica_cal = ?, freq_cal = ?, tiempo_cal = ?,
                aplica_val = ?, freq_val = ?, tiempo_val = ?
            WHERE id = ?
            """,
            (
                data["aplica_mp"],
                data["freq_mp"],
                data["tiempo_mp"],
                data["aplica_cal"],
                data["freq_cal"],
                data["tiempo_cal"],
                data["aplica_val"],
                data["freq_val"],
                data["tiempo_val"],
                inventario_equipo_id,
            ),
        )

        raw_params = _ensure_params(conn, inv["empresa_id"])
        params_calc = calc_parametros(raw_params)
        conn.commit()

    # Catálogo aparte: reemplaza últimos valores por descripción de equipo.
    catalogo = None
    try:
        catalogo = upsert_parametros(inv["equipo"] or "", data)
    except ValueError:
        catalogo = None

    calc = calc_fila_equipo(
        {
            **data,
            "forma_adquisicion": inv["forma_adquisicion"],
            "estado": inv["estado"] if "estado" in inv.keys() else None,
            "fecha_baja": inv["fecha_baja"] if "fecha_baja" in inv.keys() else None,
            "ubicacion": inv["ubicacion"] if "ubicacion" in inv.keys() else None,
            "codigo_ubicacion": inv["codigo_ubicacion"] if "codigo_ubicacion" in inv.keys() else None,
            "ID_servicio": inv["ID_servicio"] if "ID_servicio" in inv.keys() else None,
            "fecha_ultimo_pm": inv["fecha_ultimo_pm"] if "fecha_ultimo_pm" in inv.keys() else None,
            "aplica_mp": bool(data["aplica_mp"]),
            "aplica_cal": bool(data["aplica_cal"]),
            "aplica_val": bool(data["aplica_val"]),
            "aplica_cap": bool(data["aplica_cap"]),
            "tercerizado_mp": bool(data["tercerizado_mp"]),
            "tercerizado_cal": bool(data["tercerizado_cal"]),
            "tercerizado_val": bool(data["tercerizado_val"]),
            "tiene_historico_correctivo": bool(data["tiene_historico_correctivo"]),
        },
        params_calc["correctivo_sin_historico"],
        anio=_anio_dim(),
    )
    meta = _audit_dimensionamiento(
        empresa_id=inv["empresa_id"],
        sede_id=inv["sede_id"],
        parametros={
            "accion": "put_equipo_actividad",
            "inventario_equipo_id": inventario_equipo_id,
            "criticidad": data["criticidad"],
        },
        resultado={"total_horas": calc.get("total_horas"), "en_plantilla": calc.get("en_plantilla")},
    )
    return jsonify(
        {
            "ok": True,
            "message": (
                "Actividad guardada. Los parámetros del tipo de equipo quedaron "
                "actualizados en el catálogo (último valor)."
            ),
            "equipo": {"inventario_equipo_id": inventario_equipo_id, **data, **calc},
            "catalogo_parametros": catalogo,
            "formula": meta,
        }
    )


@dim_bp.post("/grupos/actividad-adicional")
@login_required
@permission_required("view_dimensionamiento")
def post_actividad_adicional_lote():
    """
    Suma una actividad adicional (h × periodicidad anual) a todas las
    unidades en dotación del lote (misma descripción de equipo).
    """
    if not _can_edit_adicionales():
        return jsonify(
            {
                "ok": False,
                "error": "Solo cargos del rol operativo pueden agregar actividades adicionales.",
            }
        ), 403

    body = request.get_json(silent=True) or {}
    try:
        sede_id = int(body.get("sede_id"))
    except (TypeError, ValueError):
        return jsonify({"ok": False, "error": "sede_id inválido."}), 400

    grupo_key = sanitize_string(body.get("grupo_key"))
    tipo = sanitize_string(body.get("tipo"))
    field = ADICIONAL_TIPOS.get(tipo)
    if not grupo_key:
        return jsonify({"ok": False, "error": "Selecciona el lote de equipo."}), 400
    if not field:
        return jsonify(
            {
                "ok": False,
                "error": "Tipo de actividad inválido. Use gestión documental, acompañamiento u otras tareas.",
            }
        ), 400

    horas = _f(body.get("horas"))
    periodicidad = _f(body.get("periodicidad_anual"))
    if horas <= 0:
        return jsonify({"ok": False, "error": "Las horas invertidas deben ser mayores a 0."}), 400
    if periodicidad <= 0:
        return jsonify(
            {"ok": False, "error": "La periodicidad anual debe ser mayor a 0."}
        ), 400

    add_h = round(horas * periodicidad, 4)
    user = current_user()
    if not can_access_sede(user, sede_id):
        return jsonify({"ok": False, "error": "No autorizado para esta sede."}), 403

    with get_empresas_connection() as conn:
        sede = _sede_meta(conn, sede_id)
        if not sede:
            return jsonify({"ok": False, "error": "Sede no encontrada."}), 404
        raw_params = _ensure_params(conn, sede["empresa_id"])
        params_calc = calc_parametros(raw_params)
        equipos = _load_equipos_contexto(conn, sede_id, params_calc)
        miembros = [
            eq
            for eq in equipos
            if eq.get("en_dotacion", True)
            and descripcion_key(eq) == grupo_key
        ]
        if not miembros:
            return jsonify(
                {"ok": False, "error": "No hay equipos en dotación para este lote."}
            ), 404

        for eq in miembros:
            eq_id = eq["inventario_equipo_id"]
            conn.execute(
                """
                INSERT OR IGNORE INTO dim_actividad_equipo (inventario_equipo_id)
                VALUES (?)
                """,
                (eq_id,),
            )
            conn.execute(
                f"""
                UPDATE dim_actividad_equipo
                SET {field} = COALESCE({field}, 0) + ?,
                    updated_at = datetime('now')
                WHERE inventario_equipo_id = ?
                """,
                (add_h, eq_id),
            )
        conn.commit()

    nombres = {
        "gestion_documental": "Gestión documental",
        "acompanamiento": "Acompañamiento",
        "otras_tareas": "Otras tareas",
    }
    lote_h = round(add_h * len(miembros), 2)
    meta = _audit_dimensionamiento(
        empresa_id=sede["empresa_id"],
        sede_id=sede_id,
        parametros={
            "accion": "actividad_adicional",
            "grupo_key": grupo_key,
            "tipo": tipo,
            "horas_unidad": add_h,
        },
        resultado={"unidades": len(miembros), "horas_anio_lote": lote_h},
    )
    return jsonify(
        {
            "ok": True,
            "message": (
                f"Se agregó {nombres[tipo]}: {add_h:g} h/año por equipo "
                f"({len(miembros)} unidad(es), {lote_h:g} h/año del lote)."
            ),
            "unidades": len(miembros),
            "horas_anio_unidad": add_h,
            "horas_anio_lote": lote_h,
            "formula": meta,
        }
    )


def _export_rows(sede_id: int):
    user = current_user()
    if not can_access_sede(user, sede_id):
        raise PermissionError("No autorizado para esta sede.")
    with get_empresas_connection() as conn:
        sede = _sede_meta(conn, sede_id)
        if not sede:
            raise ValueError("Sede no encontrada.")
        raw_params = _ensure_params(conn, sede["empresa_id"])
        params_calc = calc_parametros(raw_params)
        equipos = _load_equipos_contexto(conn, sede_id, params_calc, anio=_anio_dim())
        grupos = build_grupos_por_descripcion(equipos)
    resumen = calc_resumen(equipos, params_calc)
    disponibles = _count_ingenieros_disponibles(sede_id, sede["empresa_id"])
    return sede, params_calc, equipos, resumen, disponibles, grupos


EXPORT_COLS = [
    "codigo",
    "servicio",
    "equipo",
    "marca",
    "modelo",
    "serial",
    "ubicacion",
                "criticidad",
    "alto_unico_area",
    "alto_especializado",
    "criticidad_plantilla",
    "aplica_mp",
    "freq_mp",
    "tiempo_mp",
    "horas_mp",
    "aplica_cal",
    "freq_cal",
    "tiempo_cal",
    "horas_cal",
    "aplica_val",
    "freq_val",
    "tiempo_val",
    "horas_val",
    "tercerizado_mp",
    "tercerizado_cal",
    "tercerizado_val",
    "aplica_cap",
    "freq_cap",
    "tiempo_cap",
    "horas_cap",
    "tiene_historico_correctivo",
    "historico_correctivo_h",
    "horas_correctivo",
    "gestion_documental_h",
    "acompanamiento_h",
    "otras_tareas_h",
    "total_horas",
    "observaciones",
]


@dim_bp.get("/sedes/<int:sede_id>/export.csv")
@login_required
@permission_required("view_dimensionamiento")
def export_csv(sede_id: int):
    try:
        sede, _params_calc, equipos, resumen, disponibles, _grupos = _export_rows(sede_id)
    except PermissionError as err:
        return jsonify({"ok": False, "error": str(err)}), 403
    except ValueError as err:
        return jsonify({"ok": False, "error": str(err)}), 404

    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=EXPORT_COLS, extrasaction="ignore")
    writer.writeheader()
    for eq in equipos:
        row = {k: eq.get(k, "") for k in EXPORT_COLS}
        for flag in (
            "aplica_mp",
            "aplica_cal",
            "aplica_val",
            "aplica_cap",
            "tiene_historico_correctivo",
            "tercerizado_mp",
            "tercerizado_cal",
            "tercerizado_val",
        ):
            row[flag] = "Sí" if eq.get(flag) else "No"
        writer.writerow(row)

    writer.writerow({})
    writer.writerow({"codigo": "RESUMEN"})
    writer.writerow({"codigo": "Total equipos", "servicio": resumen["total_equipos"]})
    writer.writerow({"codigo": "Total horas base", "servicio": resumen["total_base"]})
    writer.writerow({"codigo": "Contingencia", "servicio": resumen["contingencia_h"]})
    writer.writerow({"codigo": "Total horas ajustadas", "servicio": resumen["total_ajustado"]})
    writer.writerow({"codigo": "Horas efectivas IC", "servicio": resumen["horas_efectivas"]})
    writer.writerow(
        {"codigo": "Ingenieros requeridos", "servicio": resumen["ingenieros_requeridos"]}
    )
    writer.writerow(
        {
            "codigo": "Ingenieros recomendados",
            "servicio": resumen["ingenieros_recomendados"],
        }
    )
    writer.writerow({"codigo": "Ingenieros disponibles", "servicio": disponibles})

    data = io.BytesIO(buffer.getvalue().encode("utf-8-sig"))
    data.seek(0)
    fname = f"dimensionamiento_{sede['ID_sede']}_{datetime.now():%Y%m%d_%H%M%S}.csv"
    return send_file(data, mimetype="text/csv", as_attachment=True, download_name=fname)


@dim_bp.get("/sedes/<int:sede_id>/export.xlsx")
@login_required
@permission_required("view_dimensionamiento")
def export_xlsx(sede_id: int):
    try:
        sede, params_calc, equipos, resumen, disponibles, _grupos = _export_rows(sede_id)
    except PermissionError as err:
        return jsonify({"ok": False, "error": str(err)}), 403
    except ValueError as err:
        return jsonify({"ok": False, "error": str(err)}), 404

    wb = Workbook()
    ws = wb.active
    ws.title = "Inventario"
    ws.append(EXPORT_COLS)
    for eq in equipos:
        row = []
        for k in EXPORT_COLS:
            val = eq.get(k, "")
            if k in (
                "aplica_mp",
                "aplica_cal",
                "aplica_val",
                "aplica_cap",
                "tiene_historico_correctivo",
                "tercerizado_mp",
                "tercerizado_cal",
                "tercerizado_val",
            ):
                val = "Sí" if eq.get(k) else "No"
            row.append(val)
        ws.append(row)

    ws_inf = wb.create_sheet("Informe_Exportable")
    ws_inf.append(["INFORME EXPORTABLE - DIMENSIONAMIENTO DE PERSONAL"])
    ws_inf.append([])
    ws_inf.append(["Sede", sede["name_sede"]])
    ws_inf.append(["Perfil", params_calc.get("perfil")])
    ws_inf.append(["Total equipos evaluados", resumen["total_equipos"]])
    ws_inf.append(["Total horas base", resumen["total_base"]])
    ws_inf.append(["Contingencia general", resumen["contingencia_h"]])
    ws_inf.append(["Total horas ajustadas", resumen["total_ajustado"]])
    ws_inf.append(["Horas efectivas por Ingeniero Clínico", resumen["horas_efectivas"]])
    ws_inf.append(["Ingenieros Clínicos requeridos", resumen["ingenieros_requeridos"]])
    ws_inf.append(["Ingenieros Clínicos recomendados", resumen["ingenieros_recomendados"]])
    ws_inf.append(["Ingenieros disponibles (personal)", disponibles])
    ws_inf.append([])
    ws_inf.append(["Resumen por servicio"])
    ws_inf.append(["Servicio", "Horas/año"])
    for item in resumen["por_servicio"]:
        ws_inf.append([item["servicio"], item["horas"]])

    ws_dash = wb.create_sheet("Dashboard")
    ws_dash.append(["Indicador", "Resultado"])
    ws_dash.append(["Total equipos", resumen["total_equipos"]])
    ws_dash.append(["Horas ajustadas/año", resumen["total_ajustado"]])
    ws_dash.append(["Horas efectivas por ingeniero", resumen["horas_efectivas"]])
    ws_dash.append(["Ingenieros requeridos", resumen["ingenieros_requeridos"]])
    ws_dash.append(["Ingenieros recomendados", resumen["ingenieros_recomendados"]])
    ws_dash.append(["Ingenieros disponibles", disponibles])

    data = io.BytesIO()
    wb.save(data)
    data.seek(0)
    fname = f"dimensionamiento_{sede['ID_sede']}_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    return send_file(
        data,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=fname,
    )


@dim_bp.get("/sedes/<int:sede_id>/export.pdf")
@login_required
@permission_required("view_dimensionamiento")
def export_pdf(sede_id: int):
    try:
        sede, params_calc, _equipos, resumen, disponibles, grupos = _export_rows(sede_id)
    except PermissionError as err:
        return jsonify({"ok": False, "error": str(err)}), 403
    except ValueError as err:
        return jsonify({"ok": False, "error": str(err)}), 404

    from server.dimensionamiento_pdf import build_informe_pdf

    sede_d = dict(sede)
    pdf_bytes = build_informe_pdf(
        sede=sede_d,
        params_calc=params_calc,
        resumen=resumen,
        disponibles=disponibles,
        grupos=grupos,
        staff=collect_org_staff(empresa_id=sede_d.get("empresa_id"), sede_id=sede_d.get("id")),
    )
    data = io.BytesIO(pdf_bytes)
    data.seek(0)
    fname = f"informe_dimensionamiento_{sede['ID_sede']}_{datetime.now():%Y%m%d_%H%M%S}.pdf"
    return send_file(data, mimetype="application/pdf", as_attachment=True, download_name=fname)
