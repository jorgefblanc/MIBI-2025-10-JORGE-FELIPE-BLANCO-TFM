"""Plantillas e informes de ejemplo para la guía de uso (se generan al vuelo)."""

from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.lib import colors

EXAMPLE_FILES = frozenset({"plantilla.xlsx", "ejemplo.pdf"})

_PLANTILLAS = {
    "suficiencia": {
        "title": "Ejemplo — Suficiencia de equipos",
        "headers": ["Servicio", "Equipo", "Serial", "Cantidad operativa", "En mantenimiento", "Disponible"],
        "rows": [
            ["Urgencias", "Monitor de signos", "MN-001", 8, 1, 7],
            ["Urgencias", "Ventilador", "VT-014", 4, 0, 4],
        ],
        "note": "Esto es un ejemplo. El inventario real se carga en el módulo Inventario.",
        "pdf_body": (
            "Este informe de ejemplo muestra cómo se lee el semáforo: "
            "Suficiente (hay equipos), Alerta (hay que vigilar) e Insuficiente (faltan equipos). "
            "En la app elige empresa, sede y servicio, y exporta tu propio Excel cuando el análisis esté listo."
        ),
        "pdf_table": [
            ["Equipo", "Resultado", "Qué hacer"],
            ["Monitor de signos", "Suficiente", "Seguir así"],
            ["Ventilador", "Alerta", "Revisar si alguno está parado"],
        ],
    },
    "dimensionamiento": {
        "title": "Ejemplo — Dimensionamiento de personal",
        "headers": ["Equipo", "MP propio (h/año)", "Otras tareas (h)", "En plantilla"],
        "rows": [
            ["Monitor de signos", 12, 4, "Sí"],
            ["Ventilador (comodato)", 0, 3, "No"],
        ],
        "note": "Ejemplo de carga de horas. En la app pulsa Cargar sede para ver tus números.",
        "pdf_body": (
            "El informe real te dice cuántos ingenieros clínicos necesita la sede, "
            "comparados con las personas que ya están asignadas. "
            "Si el número sale alto, no bajes la criticidad a mano: coméntalo con tu coordinador."
        ),
        "pdf_table": [
            ["Indicador", "Ejemplo"],
            ["Horas al año", "1 840"],
            ["Ingenieros necesarios", "1,2"],
            ["Personal en la sede", "1"],
        ],
    },
    "frecuencia-pm": {
        "title": "Ejemplo — Frecuencia de mantenimiento",
        "headers": ["Equipo", "Grupo", "Frecuencia (meses)", "Próxima visita"],
        "rows": [
            ["Monitor de signos", "GE-II", 6, "2026-09-15"],
            ["Ventilador", "GE-III", 3, "2026-07-01"],
        ],
        "note": "Ejemplo de cronograma. En la app sincroniza el inventario y pulsa Recalcular.",
        "pdf_body": (
            "El cronograma reparte las visitas del año. "
            "Si un mes se ve cargado, no lo cambies a mano: ajusta la frecuencia y vuelve a calcular. "
            "Puedes bajar el PDF del informe y el del calendario desde la barra del módulo."
        ),
        "pdf_table": [
            ["Mes", "Visitas de ejemplo"],
            ["Marzo", "4"],
            ["Junio", "6"],
            ["Septiembre", "4"],
        ],
    },
    "preinstalacion": {
        "title": "Ejemplo — Visita de preinstalación",
        "headers": ["Ítem", "Estado", "Comentario"],
        "rows": [
            ["Toma eléctrica dedicada", "Cumple", "Punto marcado en plano"],
            ["Ancho de puerta", "No cumple", "Faltan 8 cm"],
            ["Aire acondicionado", "Pendiente", "Medición el jueves"],
        ],
        "note": "Ejemplo de visita SITIO. En la app crea una Nueva visita y recorre las exigencias.",
        "pdf_body": (
            "El PDF de la visita deja constancia de lo que cumple el área y lo que falta. "
            "Un Pendiente se trata como no listo para instalar. "
            "Adjunta foto cuando algo estructural o eléctrico no cumpla."
        ),
        "pdf_table": [
            ["Resumen", "Cantidad"],
            ["Cumple", "18"],
            ["Pendiente", "2"],
            ["No cumple", "1"],
        ],
    },
    "kpis": {
        "title": "Ejemplo — KPIs del mes",
        "headers": ["Indicador", "Resultado", "Semáforo"],
        "rows": [
            ["Cumplimiento MP", "95 %", "Verde"],
            ["Disponibilidad", "92 %", "Amarillo"],
            ["Días fuera de servicio", "18", "Rojo"],
        ],
        "note": "Ejemplo gerencial. En la app carga empresa y año, completa el mes y exporta Excel.",
        "pdf_body": (
            "Verde: vas bien. Amarillo: vigila. Rojo: hay que anotar una acción en el plan "
            "(quién, qué y para cuándo). No cambies la meta a mitad de año solo para pintar verde."
        ),
        "pdf_table": [
            ["Mes", "Semáforo general"],
            ["Enero", "Amarillo"],
            ["Febrero", "Verde"],
            ["Marzo", "Rojo"],
        ],
    },
}


def _spec(module: str) -> dict:
    return _PLANTILLAS.get(module) or _PLANTILLAS["suficiencia"]


def build_plantilla_xlsx(module: str) -> io.BytesIO:
    spec = _spec(module)
    wb = Workbook()
    ws = wb.active
    ws.title = "Ejemplo"
    ws["A1"] = spec["title"]
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = spec["note"]
    headers = spec["headers"]
    for col, name in enumerate(headers, start=1):
        cell = ws.cell(4, col, name)
        cell.font = Font(bold=True)
    for r, row in enumerate(spec["rows"], start=5):
        for c, value in enumerate(row, start=1):
            ws.cell(r, c, value)
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[chr(64 + col)].width = 22
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def build_ejemplo_pdf(module: str) -> io.BytesIO:
    spec = _spec(module)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=2 * cm,
        rightMargin=2 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
        title=spec["title"],
    )
    styles = getSampleStyleSheet()
    title = ParagraphStyle(
        "help_ex_title",
        parent=styles["Heading1"],
        fontSize=16,
        textColor=colors.HexColor("#111827"),
        spaceAfter=8,
    )
    body = ParagraphStyle(
        "help_ex_body",
        parent=styles["BodyText"],
        fontSize=11,
        leading=15,
        textColor=colors.HexColor("#111827"),
    )
    story = [
        Paragraph("SIGTB · Ejemplo de informe (no es un dato real)", styles["Normal"]),
        Spacer(1, 6),
        Paragraph(spec["title"], title),
        Paragraph(spec["pdf_body"], body),
        Spacer(1, 12),
    ]
    table = Table(spec["pdf_table"], colWidths=[8 * cm, 8 * cm])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DBEAFE")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 10),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D1D5DB")),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    story.append(table)
    story.append(Spacer(1, 14))
    story.append(
        Paragraph(
            f"Generado como ejemplo el {datetime.now():%Y-%m-%d}. Usa los botones de exportar del módulo para tu informe real.",
            body,
        )
    )
    doc.build(story)
    buf.seek(0)
    return buf


def build_example(module: str, filename: str) -> tuple[io.BytesIO, str, str]:
    if filename == "plantilla.xlsx":
        return (
            build_plantilla_xlsx(module),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            f"ejemplo_{module}_plantilla.xlsx",
        )
    return (
        build_ejemplo_pdf(module),
        "application/pdf",
        f"ejemplo_{module}_informe.pdf",
    )
