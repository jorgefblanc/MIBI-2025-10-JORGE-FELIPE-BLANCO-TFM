"""API HTTP de KPIs de ingeniería clínica."""

from __future__ import annotations

import io
import json
from datetime import datetime

from flask import Blueprint, jsonify, request, send_file
from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Font

from server.authz import (
    assigned_empresa_id,
    can_access_empresa,
    current_user,
    is_global_scope,
    login_required,
    permission_required,
)
from server.db import get_empresas_connection, get_users_connection
from server.event_log import log_evento
from server.kpis import (
    ESTADOS_DECISION,
    INDICATORS,
    MESES,
    build_year,
    catalogo_payload,
    compute_year,
    ensure_kpis_schema,
    merge_params,
    sanitize_decision,
    sanitize_month_payload,
    sanitize_params_payload,
    snapshot_ejecucion_por_mes,
    snapshot_inventario,
)
from server.limits import JSON_SMALL

kpis_bp = Blueprint("kpis", __name__, url_prefix="/api/kpis")
KPI_PERMS = ("view_kpis", "view_respaldo")
MAX_REQUEST_BYTES = JSON_SMALL * 2 + 8_000


def _reject_huge_body():
    length = request.content_length
    if length is not None and length > MAX_REQUEST_BYTES:
        return jsonify({"ok": False, "error": "La solicitud es demasiado grande."}), 413
    return None


def _json_obj(raw) -> dict:
    if not raw:
        return {}
    if isinstance(raw, dict):
        return raw
    try:
        data = json.loads(raw)
    except (TypeError, ValueError, json.JSONDecodeError):
        return {}
    return data if isinstance(data, dict) else {}


def _resolve_empresa(user, requested) -> tuple[int | None, object | None]:
    try:
        empresa_id = int(requested) if requested not in (None, "", "null") else None
    except (TypeError, ValueError):
        return None, (jsonify({"ok": False, "error": "empresa_id inválido."}), 400)
    if is_global_scope(user):
        if not empresa_id:
            return None, (jsonify({"ok": False, "error": "Selecciona una empresa."}), 400)
    else:
        assigned = assigned_empresa_id(user)
        if not assigned:
            return None, (jsonify({"ok": False, "error": "Sin empresa asignada."}), 403)
        if empresa_id and empresa_id != assigned:
            return None, (jsonify({"ok": False, "error": "Solo puedes ver KPIs de tu empresa."}), 403)
        empresa_id = assigned
    with get_empresas_connection() as emp_conn:
        if not can_access_empresa(emp_conn, user, empresa_id):
            return None, (jsonify({"ok": False, "error": "Fuera de tu empresa."}), 403)
        row = emp_conn.execute(
            "SELECT id, ID_Empresa FROM empresas WHERE id = ?",
            (empresa_id,),
        ).fetchone()
    if not row:
        return None, (jsonify({"ok": False, "error": "La empresa no existe."}), 404)
    return empresa_id, None


def _load_params(emp_conn, empresa_id: int, empresa_nombre: str) -> dict:
    row = emp_conn.execute(
        "SELECT * FROM kpis_parametros WHERE empresa_id = ?",
        (empresa_id,),
    ).fetchone()
    stored = None
    if row:
        stored = {
            "anio": row["anio"],
            "institucion": row["institucion"],
            "proceso": row["proceso"],
            "responsable": row["responsable"],
            "costo_diario_inactividad": row["costo_diario_inactividad"],
            "factor_mitigacion": row["factor_mitigacion"],
            "horas_laborales_tecnico_mes": row["horas_laborales_tecnico_mes"],
            "costo_externo_referencia": row["costo_externo_referencia"],
            "metas": _json_obj(row["metas_json"]),
        }
    params = merge_params(stored, empresa_nombre)
    if not params.get("institucion"):
        params["institucion"] = empresa_nombre
    return params


def _load_months(emp_conn, empresa_id: int, anio: int) -> dict[int, dict]:
    rows = emp_conn.execute(
        """
        SELECT mes, payload_json FROM kpis_entrada_mes
        WHERE empresa_id = ? AND anio = ?
        """,
        (empresa_id, anio),
    ).fetchall()
    return {int(r["mes"]): _json_obj(r["payload_json"]) for r in rows}


def _load_decisiones(emp_conn, empresa_id: int, anio: int) -> dict[str, dict]:
    rows = emp_conn.execute(
        """
        SELECT indicador_key, estado, observaciones
        FROM kpis_decisiones
        WHERE empresa_id = ? AND anio = ?
        """,
        (empresa_id, anio),
    ).fetchall()
    return {
        r["indicador_key"]: {"estado": r["estado"], "observaciones": r["observaciones"] or ""}
        for r in rows
    }


def _assemble(emp_conn, users_conn, _user, empresa_id: int, anio: int | None):
    ensure_kpis_schema(emp_conn)
    emp = emp_conn.execute(
        "SELECT id, ID_Empresa FROM empresas WHERE id = ?",
        (empresa_id,),
    ).fetchone()
    nombre = emp["ID_Empresa"] if emp else f"Empresa {empresa_id}"
    params = _load_params(emp_conn, empresa_id, nombre)
    year = int(anio or params.get("anio") or datetime.now().year)
    year = max(2000, min(2100, year))
    months = build_year(year, params, _load_months(emp_conn, empresa_id, year))
    anual = compute_year(months, params)
    decisiones = _load_decisiones(emp_conn, empresa_id, year)
    plan = []
    for row in anual["indicadores"]:
        extra = decisiones.get(row["key"]) or {}
        estado = extra.get("estado") or "Pendiente"
        plazo = "30 días" if row["prioridad"] == "Alta" else ("60 días" if row["prioridad"] == "Media" else "Seguimiento rutinario")
        if row["semaforo"] == "Verde":
            accion = "Mantener seguimiento"
        elif row["semaforo"] == "Amarillo":
            accion = "Ejecutar mejora controlada"
        else:
            accion = "Intervención prioritaria y reporte a gerencia"
        plan.append(
            {
                **row,
                "estado": estado if estado in ESTADOS_DECISION else "Pendiente",
                "observaciones": extra.get("observaciones") or "",
                "plazo": plazo,
                "responsable": "Ingeniería Clínica",
                "accion_plan": accion,
            }
        )
    try:
        snapshot = snapshot_inventario(emp_conn, users_conn, empresa_id, year)
    except Exception:
        snapshot = {
            "equipos_totales": 0,
            "equipos_criticos": 0,
            "tecnicos": 0,
            "fuera_servicio_unidades": 0,
            "mp_programados": 0,
            "cal_programadas": 0,
            "equipos_documentacion_completa": 0,
            "equipos_fuera_soporte": 0,
            "equipos_candidatos_capex": 0,
            "equipos_bodega_baja": 0,
            "equipos_sin_datos_anio": 0,
            "nota": "No se pudo leer el inventario. Complete los conteos de parque a mano.",
        }
    try:
        snapshot["ejecucion_mes"] = snapshot_ejecucion_por_mes(emp_conn, empresa_id, year)
    except Exception:
        snapshot["ejecucion_mes"] = {}
    return {
        "empresa_id": empresa_id,
        "empresa": nombre,
        "anio": year,
        "params": params,
        "meses": months,
        "anual": anual,
        "plan": plan,
        "snapshot": snapshot,
        "catalogo": catalogo_payload(),
        "can_edit": True,
    }


@kpis_bp.get("/contexto")
@login_required
@permission_required(*KPI_PERMS)
def get_contexto():
    user = current_user()
    empresa_id, err = _resolve_empresa(user, request.args.get("empresa_id"))
    if err:
        return err
    try:
        anio = int(request.args.get("anio") or 0) or None
    except (TypeError, ValueError):
        return jsonify({"ok": False, "error": "Año inválido."}), 400
    with get_empresas_connection() as emp_conn, get_users_connection() as users_conn:
        data = _assemble(emp_conn, users_conn, user, empresa_id, anio)
    return jsonify({"ok": True, **data})


@kpis_bp.route("/parametros", methods=["GET", "PUT"])
@login_required
@permission_required(*KPI_PERMS)
def parametros():
    """GET lee parámetros; PUT los guarda (misma URL para el cliente)."""
    if request.method == "GET":
        return _get_parametros()
    return _put_parametros()


def _get_parametros():
    user = current_user()
    empresa_id, err = _resolve_empresa(user, request.args.get("empresa_id"))
    if err:
        return err
    with get_empresas_connection() as emp_conn:
        ensure_kpis_schema(emp_conn)
        emp = emp_conn.execute(
            "SELECT id, ID_Empresa FROM empresas WHERE id = ?",
            (empresa_id,),
        ).fetchone()
        nombre = emp["ID_Empresa"] if emp else f"Empresa {empresa_id}"
        params = _load_params(emp_conn, empresa_id, nombre)
    return jsonify({"ok": True, "empresa_id": empresa_id, "params": params})


def _put_parametros():
    huge = _reject_huge_body()
    if huge:
        return huge
    user = current_user()
    body = request.get_json(silent=True) or {}
    empresa_id, err = _resolve_empresa(user, body.get("empresa_id"))
    if err:
        return err
    try:
        cleaned = sanitize_params_payload(body)
    except ValueError as exc:
        return jsonify({"ok": False, "error": str(exc)}), 400
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    with get_empresas_connection() as emp_conn:
        ensure_kpis_schema(emp_conn)
        emp_conn.execute(
            """
            INSERT INTO kpis_parametros (
                empresa_id, anio, institucion, proceso, responsable,
                costo_diario_inactividad, factor_mitigacion,
                horas_laborales_tecnico_mes, costo_externo_referencia,
                metas_json, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(empresa_id) DO UPDATE SET
                anio = excluded.anio,
                institucion = excluded.institucion,
                proceso = excluded.proceso,
                responsable = excluded.responsable,
                costo_diario_inactividad = excluded.costo_diario_inactividad,
                factor_mitigacion = excluded.factor_mitigacion,
                horas_laborales_tecnico_mes = excluded.horas_laborales_tecnico_mes,
                costo_externo_referencia = excluded.costo_externo_referencia,
                metas_json = excluded.metas_json,
                updated_at = excluded.updated_at
            """,
            (
                empresa_id,
                cleaned["anio"],
                cleaned["institucion"],
                cleaned["proceso"],
                cleaned["responsable"],
                cleaned["costo_diario_inactividad"],
                cleaned["factor_mitigacion"],
                cleaned["horas_laborales_tecnico_mes"],
                cleaned["costo_externo_referencia"],
                json.dumps(cleaned["metas"], ensure_ascii=False),
                now,
            ),
        )
        emp_conn.commit()
    log_evento(
        empresa_id=empresa_id,
        accion=f"Actualizó parámetros de KPIs (año {cleaned['anio']}).",
        user=user,
    )
    return jsonify({"ok": True, "message": "Parámetros guardados.", "params": cleaned})


@kpis_bp.put("/entrada/<int:mes>")
@login_required
@permission_required(*KPI_PERMS)
def put_entrada(mes: int):
    huge = _reject_huge_body()
    if huge:
        return huge
    user = current_user()
    if mes < 1 or mes > 12:
        return jsonify({"ok": False, "error": "Mes inválido."}), 400
    body = request.get_json(silent=True) or {}
    empresa_id, err = _resolve_empresa(user, body.get("empresa_id"))
    if err:
        return err
    try:
        anio = int(body.get("anio") or datetime.now().year)
    except (TypeError, ValueError):
        return jsonify({"ok": False, "error": "Año inválido."}), 400
    anio = max(2000, min(2100, anio))
    try:
        cleaned = sanitize_month_payload(body.get("entrada") or body)
    except ValueError as exc:
        return jsonify({"ok": False, "error": str(exc)}), 400
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    uid = int(user.get("id_usuario") or 0) or None
    with get_empresas_connection() as emp_conn:
        ensure_kpis_schema(emp_conn)
        emp_conn.execute(
            """
            INSERT INTO kpis_entrada_mes (
                empresa_id, anio, mes, payload_json, updated_at, updated_by
            ) VALUES (?, ?, ?, ?, ?, ?)
            ON CONFLICT(empresa_id, anio, mes) DO UPDATE SET
                payload_json = excluded.payload_json,
                updated_at = excluded.updated_at,
                updated_by = excluded.updated_by
            """,
            (
                empresa_id,
                anio,
                mes,
                json.dumps(cleaned, ensure_ascii=False),
                now,
                uid,
            ),
        )
        emp_conn.commit()
    log_evento(
        empresa_id=empresa_id,
        accion=f"Guardó entrada de KPIs {MESES[mes - 1]} {anio}.",
        user=user,
    )
    return jsonify({"ok": True, "message": f"Entrada de {MESES[mes - 1]} guardada.", "mes": mes})


@kpis_bp.put("/decisiones/<indicador_key>")
@login_required
@permission_required(*KPI_PERMS)
def put_decision(indicador_key: str):
    huge = _reject_huge_body()
    if huge:
        return huge
    user = current_user()
    if indicador_key not in {i["key"] for i in INDICATORS}:
        return jsonify({"ok": False, "error": "Indicador desconocido."}), 400
    body = request.get_json(silent=True) or {}
    empresa_id, err = _resolve_empresa(user, body.get("empresa_id"))
    if err:
        return err
    try:
        anio = int(body.get("anio") or datetime.now().year)
    except (TypeError, ValueError):
        return jsonify({"ok": False, "error": "Año inválido."}), 400
    anio = max(2000, min(2100, anio))
    estado, obs = sanitize_decision(body.get("estado"), body.get("observaciones"))
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    with get_empresas_connection() as emp_conn:
        ensure_kpis_schema(emp_conn)
        emp_conn.execute(
            """
            INSERT INTO kpis_decisiones (
                empresa_id, anio, indicador_key, estado, observaciones, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?)
            ON CONFLICT(empresa_id, anio, indicador_key) DO UPDATE SET
                estado = excluded.estado,
                observaciones = excluded.observaciones,
                updated_at = excluded.updated_at
            """,
            (empresa_id, anio, indicador_key, estado, obs, now),
        )
        emp_conn.commit()
    return jsonify({"ok": True, "estado": estado, "observaciones": obs})


@kpis_bp.get("/informe.xlsx")
@login_required
@permission_required(*KPI_PERMS)
def export_xlsx():
    user = current_user()
    empresa_id, err = _resolve_empresa(user, request.args.get("empresa_id"))
    if err:
        return err
    try:
        anio = int(request.args.get("anio") or 0) or None
    except (TypeError, ValueError):
        return jsonify({"ok": False, "error": "Año inválido."}), 400
    with get_empresas_connection() as emp_conn, get_users_connection() as users_conn:
        data = _assemble(emp_conn, users_conn, user, empresa_id, anio)
    wb = Workbook()
    inf = wb.active
    inf.title = "Informe_Alta_Gerencia"
    inf.append(["INFORME EXPORTABLE PARA ALTA GERENCIA - KPIs INGENIERÍA CLÍNICA"])
    inf.append([])
    p = data["params"]
    inf.append(["Institución", p.get("institucion") or data["empresa"]])
    inf.append(["Proceso", p.get("proceso")])
    inf.append(["Responsable", p.get("responsable")])
    inf.append(["Año", data["anio"]])
    inf.append(["Fecha generación", datetime.now().strftime("%Y-%m-%d %H:%M")])
    inf.append(["Índice gerencial promedio", data["anual"]["indice_gerencial"]])
    inf.append([])
    inf.append(["Resumen ejecutivo", data["anual"]["resumen_ejecutivo"]])
    inf.append(["Decisión sugerida", data["anual"]["decision_ejecutiva"]])
    inf.append([])
    inf.append(["Tipo", "Indicador", "Resultado", "Semáforo", "Interpretación", "Mensaje gerencial"])
    for row in data["anual"]["indicadores"]:
        inf.append(
            [
                row["tipo"],
                row["nombre"],
                row["resultado"],
                row["semaforo"],
                row["analisis"],
                row["mensaje_gerencial"],
            ]
        )
    inf["A1"].font = Font(bold=True)

    dash = wb.create_sheet("Dashboard")
    dash.append(["Bloque", "Indicador clave", "Resultado", "Meta", "Semáforo", "Lectura"])
    for row in data["anual"]["dashboard"]:
        dash.append([row["tipo"], row["nombre"], row["resultado"], row["meta"], row["semaforo"], row["analisis"]])

    sem_counts = {"Verde": 0, "Amarillo": 0, "Rojo": 0}
    for row in data["anual"]["dashboard"]:
        sem = row.get("semaforo")
        if sem in sem_counts:
            sem_counts[sem] += 1
    dash.append([])
    dash.append(["Semáforo", "Indicadores"])
    pie_start = dash.max_row + 1
    for label, n in sem_counts.items():
        dash.append([label, n])
    pie = PieChart()
    pie.title = "Criticidad (semáforo)"
    labels = Reference(dash, min_col=1, min_row=pie_start, max_row=pie_start + 2)
    data_ref = Reference(dash, min_col=2, min_row=pie_start - 1, max_row=pie_start + 2)
    pie.add_data(data_ref, titles_from_data=True)
    pie.set_categories(labels)
    pie.dataLabels = DataLabelList()
    pie.dataLabels.showPercent = True
    pie.dataLabels.showVal = False
    pie.dataLabels.showCatName = True
    pie.width = 12
    pie.height = 8
    dash.add_chart(pie, "H2")

    dash.append([])
    dash.append(["Indicador", "Resultado"])
    bar_start = dash.max_row + 1
    for row in data["anual"]["dashboard"]:
        try:
            val = float(row.get("resultado") or 0)
        except (TypeError, ValueError):
            val = 0
        dash.append([row.get("nombre") or "—", val])
    bar_end = dash.max_row
    if bar_end >= bar_start:
        bar = BarChart()
        bar.type = "col"
        bar.title = "Resultado de indicadores"
        bar.y_axis.title = "Valor"
        bar_data = Reference(dash, min_col=2, min_row=bar_start - 1, max_row=bar_end)
        bar_cats = Reference(dash, min_col=1, min_row=bar_start, max_row=bar_end)
        bar.add_data(bar_data, titles_from_data=True)
        bar.set_categories(bar_cats)
        bar.shape = 4
        bar.width = 16
        bar.height = 8
        dash.add_chart(bar, "H18")

    anual = wb.create_sheet("Analisis_Anual")
    anual.append(["Tipo", "Indicador", "Resultado anual", "Meta", "Unidad", "Semáforo", "Prioridad", "Análisis", "Decisión"])
    for row in data["anual"]["indicadores"]:
        anual.append(
            [
                row["tipo"],
                row["nombre"],
                row["resultado"],
                row["meta"],
                row["unidad"],
                row["semaforo"],
                row["prioridad"],
                row["analisis"],
                row["decision"],
            ]
        )

    mens = wb.create_sheet("KPIs_Mensuales")
    headers = ["Mes"] + [i["nombre"] for i in INDICATORS] + ["Índice gerencial", "Semáforo global"]
    mens.append(headers)
    for month in data["meses"]:
        line = [month["nombre"]]
        for ind in INDICATORS:
            line.append(month["kpis"].get(ind["key"]))
        line.append(month["kpis"].get("indice_gerencial"))
        line.append(month["kpis"].get("semaforo_global"))
        mens.append(line)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"kpis_{data['empresa_id']}_{data['anio']}.xlsx"
    return send_file(
        buf,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
